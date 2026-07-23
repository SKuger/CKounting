"""Parseo de la respuesta de MUISCA (empresa vs persona natural) + export.
No usa red: prueba `_parse_nombre` con fragmentos HTML como los que devuelve
`DefSelPublicacionesExterna.faces`."""
from pathlib import Path

from openpyxl import load_workbook

from contab_core import dian_inconsistencias as di

# Empresa (persona jurídica): la celda 'Razón Social'.
EMPRESA_HTML = """
<tr><td class="lbl">Razón Social</td><td class="val">SEGUROS COMERCIALES BOLIVAR S.A.</td></tr>
"""

# Persona natural: 4 campos (dos por fila, como en el portal).
PERSONA_HTML = """
<tr><td>Primer Apellido</td><td>PATIÑO</td><td>Segundo Apellido</td><td>CHICA</td></tr>
<tr><td>Primer Nombre</td><td>ALDEMAR</td><td>Otros Nombres</td><td></td></tr>
"""

PERSONA_2NOMBRES_HTML = """
<tr><td>Primer Apellido</td><td>KUGER</td><td>Segundo Apellido</td><td>ARBELAEZ</td></tr>
<tr><td>Primer Nombre</td><td>CARLOS</td><td>Otros Nombres</td><td>ALBERTO</td></tr>
"""

NADA_HTML = "<tr><td>Fecha Actual de Consulta</td><td>23-07-2026</td></tr>"


def test_parse_empresa():
    assert di._parse_nombre(EMPRESA_HTML) == ("SEGUROS COMERCIALES BOLIVAR S.A.", "empresa")


def test_parse_persona_un_nombre():
    # Orden natural: nombres + apellidos.
    assert di._parse_nombre(PERSONA_HTML) == ("ALDEMAR PATIÑO CHICA", "persona")


def test_parse_persona_dos_nombres():
    assert di._parse_nombre(PERSONA_2NOMBRES_HTML) == ("CARLOS ALBERTO KUGER ARBELAEZ", "persona")


def test_parse_sin_resultado():
    assert di._parse_nombre(NADA_HTML) == ("", None)


def test_export_incluye_tipo(tmp_path: Path):
    rows = [
        {"nit": "860002180", "dv": "7", "razon_social": "SEGUROS COMERCIALES BOLIVAR S.A.", "tipo": "empresa"},
        {"nit": "3358520", "dv": "7", "razon_social": "ALDEMAR PATIÑO CHICA", "tipo": "persona"},
    ]
    out = tmp_path / "inc.xlsx"
    res = di.export_inconsistencias(rows, out)
    assert res["rows_written"] == 2
    wb = load_workbook(out, read_only=True)
    ws = wb.active
    data = list(ws.iter_rows(values_only=True))
    wb.close()
    assert data[0] == ("NIT", "DV", "Razón Social / Nombre", "Tipo")
    assert data[1] == ("860002180", "7", "SEGUROS COMERCIALES BOLIVAR S.A.", "Empresa")
    assert data[2] == ("3358520", "7", "ALDEMAR PATIÑO CHICA", "Persona")
