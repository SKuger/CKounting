"""La bandeja que se edita, y el IVA por tarifa (observaciones del contador, 2026-10-28).

Lo que pidió:

1. Elegir el tercero de la línea de proveedores o clientes, u otro NIT.
2. Guardar el movimiento antes de aprobarlo, y aprobar uno a uno o en bloque.
3. Cuando el IVA no cuadra con la base, ver cada tarifa: parte del gasto sin
   IVA va a su propia cuenta, y una factura puede traer varias tarifas (D1).
4. Cada tarifa tiene su cuenta; las notas van a las de devolución.

Los NIT de estas pruebas son inventados.
"""
from decimal import Decimal
from pathlib import Path

import pytest

from contab_core import (
    asiento,
    causacion,
    documentos,
    perfil_fiscal,
    siigo,
    storage,
    tarifas,
    xml_ubl,
)

MUESTRAS = Path(__file__).resolve().parents[3] / "docs" / "samples" / "xml"
CUFE = "1" * 92 + "aaaa"
OTRO = "900999111"          # un tercero inventado, distinto del emisor


@pytest.fixture
def conn(tmp_path: Path):
    db = tmp_path / "e.db"
    # El NIT es el receptor de la factura de muestra (docs/samples/xml).
    storage.init_empresa_db(db, nit="900181831", dv="4", razon_social="PRUEBA")
    c = storage.connect(db)
    storage.ensure_schema(c)
    perfil_fiscal.guardar_perfil(c, {
        "vigencia_desde": "2026-01-01", "regimen": "ordinario", "responsable_iva": 1,
        "periodicidad_iva": "bimestral", "agente_retefuente": 1, "agente_reteiva": 1})
    tarifas.guardar_uvt(c, 2026, "52374")
    tarifas.guardar_concepto(c, {
        "concepto": "Compras.", "aplica_a": "ambos", "sobre": "base",
        "tarifa": "2.5", "base_minima_uvt": "10", "cuenta_puc": "23654001",
        "vigencia_desde": "2026-01-01"})
    # 1.500.000 + IVA 285.000 = 1.785.000
    lectura = xml_ubl.leer(MUESTRAS / "factura-con-iva.xml")
    documentos.guardar(c, lectura.documento, MUESTRAS / "factura-con-iva.xml")
    c.execute("INSERT INTO excel_load(id, source_path, source_name, total_rows, periodo)"
              " VALUES (1,'x.xlsx','x.xlsx',1,'2026-09')")
    c.execute("INSERT INTO excel_row(load_id,row_index,cufe_cude,data_json,periodo)"
              " VALUES (1,0,?,'{}','2026-09')", (CUFE,))
    c.execute("UPDATE empresa_info SET cuenta_proveedores = '220505',"
              " cuenta_iva_descontable = '240810'")
    concepto = tarifas.listar_conceptos(c)[0]["id"]
    causacion.guardar(c, {"tercero_nit": "800100200", "cuenta_puc": "519530",
                          "concepto_retencion_id": concepto})
    c.commit()
    yield c
    c.close()


def _uno(conn):
    return asiento.listar(conn, "2026-09")[0]


def _cuadra(lineas):
    return (sum(Decimal(x["debito"]) for x in lineas)
            == sum(Decimal(x["credito"]) for x in lineas))


def _como_d1(conn):
    """La misma factura, partida como una de D1: 1.000.000 al 19 %, 300.000 al
    5 % y 200.000 sin IVA. Subtotal 1.500.000, IVA 205.000, total 1.705.000."""
    doc = conn.execute("SELECT id FROM documento").fetchone()["id"]
    conn.execute("DELETE FROM documento_impuesto WHERE documento_id = ? AND linea_id IS NULL"
                 " AND es_retencion = 0", (doc,))
    for base, tarifa, valor in (("1000000", "19.00", "190000"), ("300000", "5.00", "15000")):
        conn.execute("INSERT INTO documento_impuesto(documento_id, tipo, nombre, base, tarifa,"
                     " valor) VALUES (?, '01', 'IVA', ?, ?, ?)", (doc, base, tarifa, valor))
    conn.execute("UPDATE documento SET subtotal = '1500000', total_impuestos = '205000',"
                 " total = '1705000' WHERE id = ?", (doc,))
    conn.commit()


# ───────────────────────────── IVA por tarifa ─────────────────────────────────

def test_una_factura_con_dos_tarifas_lleva_una_linea_de_iva_por_tarifa(conn):
    _como_d1(conn)
    asiento.generar(conn, "2026-09")
    lineas = _uno(conn)["lineas"]
    ivas = [x for x in lineas if x["rol"] == "iva"]
    assert sorted(Decimal(x["debito"]) for x in ivas) == [Decimal("15000"), Decimal("190000")]
    assert sorted(Decimal(x["base"]) for x in ivas) == [Decimal("300000"), Decimal("1000000")]
    assert _cuadra(lineas)


def test_la_parte_sin_iva_va_en_su_propia_linea_y_pide_cuenta(conn):
    _como_d1(conn)
    asiento.generar(conn, "2026-09")
    a = _uno(conn)
    gasto = next(x for x in a["lineas"] if x["rol"] == "gasto")
    sin_iva = next(x for x in a["lineas"] if x["rol"] == "no_gravado")
    assert Decimal(gasto["debito"]) == Decimal("1300000")
    assert Decimal(sin_iva["debito"]) == Decimal("200000")
    assert sin_iva["cuenta_puc"] is None
    assert a["estado"] == asiento.ESTADO_REVISAR, "una línea sin cuenta no se aprueba en lote"
    assert any("no llevan IVA" in x for x in a["avisos"])


def test_la_cuenta_de_la_parte_sin_iva_se_aprende_al_aprobar(conn):
    _como_d1(conn)
    asiento.generar(conn, "2026-09")
    a = _uno(conn)
    lineas = [dict(x) for x in a["lineas"]]
    for x in lineas:
        if x["rol"] == "no_gravado":
            x["cuenta_puc"] = "51953501"
    asiento.aprobar(conn, a["id"], lineas=lineas)
    regla = causacion.proponer(conn, {"tercero_nit": "800100200", "naturaleza": "compra"})
    assert regla.cuenta_no_gravada == "51953501"

    # El mes siguiente la cuenta se propone, pero el asiento queda para que él
    # la mire: «depende de la factura» (el contador, 2026-10-29).
    asiento.desaprobar(conn, a["id"])
    asiento.descartar(conn, a["id"])
    a = _uno(conn)
    sin_iva = next(x for x in a["lineas"] if x["rol"] == "no_gravado")
    assert sin_iva["cuenta_puc"] == "51953501"
    assert a["estado"] == asiento.ESTADO_REVISAR
    assert any("depende de la factura" in x for x in a["avisos"])


def test_cada_tarifa_toma_su_codigo_y_su_cuenta_del_catalogo(conn):
    _como_d1(conn)
    conn.executemany(
        "INSERT INTO siigo_impuesto(codigo, contexto, nombre, tipo, tarifa, cuenta)"
        " VALUES (?,?,?,?,?,?)",
        [("02", "compra", "IVA 19% compras", "iva", "19", "24081001"),
         ("03", "compra", "IVA 5% compras", "iva", "5", "24081005")])
    siigo.guardar_config(conn, {"iva:19": "02", "iva:5": "03"})
    asiento.generar(conn, "2026-09")
    ivas = {x["codigo_siigo"]: x for x in _uno(conn)["lineas"] if x["rol"] == "iva"}
    assert ivas["02"]["cuenta_puc"] == "24081001"
    assert ivas["03"]["cuenta_puc"] == "24081005"


def test_el_impuesto_al_consumo_es_mayor_valor_del_gasto(conn):
    """Un restaurante cobra INC, no IVA: no se descuenta y no entra al ReteIVA."""
    doc = conn.execute("SELECT id FROM documento").fetchone()["id"]
    conn.execute("DELETE FROM documento_impuesto WHERE documento_id = ? AND linea_id IS NULL"
                 " AND es_retencion = 0", (doc,))
    conn.execute("INSERT INTO documento_impuesto(documento_id, tipo, nombre, base, tarifa, valor)"
                 " VALUES (?, '04', 'INC', '1500000', '8.00', '120000')", (doc,))
    conn.execute("UPDATE documento SET total_impuestos = '120000', total = '1620000'"
                 " WHERE id = ?", (doc,))
    conn.commit()
    asiento.generar(conn, "2026-09")
    lineas = _uno(conn)["lineas"]
    assert not [x for x in lineas if x["rol"] == "iva"]
    gastos = [x for x in lineas if x["rol"] == "gasto"]
    assert sum(Decimal(x["debito"]) for x in gastos) == Decimal("1620000")
    assert _cuadra(lineas)


def test_sin_desglose_sale_como_antes(conn):
    """El reporte de la DIAN no trae tarifa: una sola línea de IVA."""
    asiento.generar(conn, "2026-09")
    lineas = _uno(conn)["lineas"]
    assert len([x for x in lineas if x["rol"] == "iva"]) == 1
    assert not [x for x in lineas if x["rol"] == "no_gravado"]


# ──────────────────────────── guardar sin aprobar ─────────────────────────────

def test_guardar_no_aprueba_y_volver_a_armar_no_lo_pisa(conn):
    asiento.generar(conn, "2026-09")
    a = _uno(conn)
    lineas = [dict(x) for x in a["lineas"]]
    lineas[0]["cuenta_puc"] = "51359501"
    g = asiento.guardar(conn, a["id"], lineas)
    assert g["editado"] == 1 and g["aprobado"] == 0

    asiento.generar(conn, "2026-09")
    assert _uno(conn)["lineas"][0]["cuenta_puc"] == "51359501"


def test_se_pueden_agregar_y_quitar_lineas(conn):
    asiento.generar(conn, "2026-09")
    a = _uno(conn)
    lineas = [dict(x) for x in a["lineas"]]
    gasto = lineas[0]
    # El gasto se reparte en dos cuentas: 1.000.000 y 500.000.
    gasto["debito"] = "1000000"
    lineas.insert(1, {"cuenta_puc": "51352001", "descripcion": "Mantenimiento",
                      "debito": "500000", "credito": "0", "rol": "gasto"})
    g = asiento.guardar(conn, a["id"], lineas)
    assert len(g["lineas"]) == len(a["lineas"]) + 1
    assert g["estado"] == asiento.ESTADO_COMPLETO

    sin_la_nueva = [x for x in g["lineas"] if x["cuenta_puc"] != "51352001"]
    g = asiento.guardar(conn, a["id"], sin_la_nueva)
    assert g["estado"] == asiento.ESTADO_REVISAR
    assert any("no cuadra por 500.000 pesos: sobra en el crédito" in x for x in g["avisos"])


def test_un_asiento_que_no_cuadra_se_guarda_pero_no_se_aprueba(conn):
    asiento.generar(conn, "2026-09")
    a = _uno(conn)
    lineas = [dict(x) for x in a["lineas"]]
    lineas[0]["debito"] = "1"
    asiento.guardar(conn, a["id"], lineas)
    with pytest.raises(ValueError, match="no cuadra"):
        asiento.aprobar(conn, a["id"])
    assert _uno(conn)["aprobado"] == 0


def test_lo_que_no_se_puede_guardar_se_rechaza_con_su_linea(conn):
    asiento.generar(conn, "2026-09")
    a = _uno(conn)
    lineas = [dict(x) for x in a["lineas"]]
    lineas[1]["cuenta_puc"] = "24-08"
    with pytest.raises(ValueError, match="Línea 2"):
        asiento.guardar(conn, a["id"], lineas)
    lineas[1]["cuenta_puc"] = "240810"
    lineas[1]["credito"] = "5"
    with pytest.raises(ValueError, match="débito o al crédito"):
        asiento.guardar(conn, a["id"], lineas)


def test_descartar_vuelve_a_la_propuesta_de_las_reglas(conn):
    asiento.generar(conn, "2026-09")
    a = _uno(conn)
    lineas = [dict(x) for x in a["lineas"]]
    lineas[0]["cuenta_puc"] = "51359501"
    asiento.guardar(conn, a["id"], lineas)
    d = asiento.descartar(conn, a["id"])
    assert d["editado"] == 0
    assert d["lineas"][0]["cuenta_puc"] == "519530"


def test_previsualizar_no_guarda(conn):
    asiento.generar(conn, "2026-09")
    a = _uno(conn)
    p = asiento.previsualizar(conn, a["id"], cuenta_puc="51352001")
    assert p["lineas"][0]["cuenta_puc"] == "51352001"
    assert _uno(conn)["lineas"][0]["cuenta_puc"] == "519530"


# ─────────────────────────── el tercero de la línea ───────────────────────────

def test_la_cartera_puede_ser_de_otro_tercero_y_asi_sale_a_siigo(conn, tmp_path):
    asiento.generar(conn, "2026-09")
    a = _uno(conn)
    lineas = [dict(x) for x in a["lineas"]]
    cartera = next(x for x in lineas if x["rol"] == "cartera")
    cartera["tercero_nit"] = f"{OTRO}-7"
    asiento.aprobar(conn, a["id"], lineas=lineas)
    guardada = next(x for x in _uno(conn)["lineas"] if x["rol"] == "cartera")
    assert guardada["tercero_nit"] == OTRO, "sin el dígito de verificación"

    siigo.guardar_config(conn, {"tipo:compra": "11", "ultimo:11": "0"})
    r = siigo.export_siigo(conn, "2026-09", tmp_path)
    from openpyxl import load_workbook
    ws = load_workbook(r["path"]).active
    cab = [c.value for c in ws[1]]
    col_tercero, col_cuenta = cab.index("Identificación tercero"), cab.index("Código cuenta contable")
    filas = list(ws.iter_rows(min_row=2, values_only=True))
    por_cuenta = {str(f[col_cuenta]): str(f[col_tercero]) for f in filas if f[col_cuenta]}
    assert por_cuenta[next(k for k in por_cuenta if k.startswith("2205"))] == OTRO
    assert por_cuenta[next(k for k in por_cuenta if k.startswith("5195"))] == "800100200"


def test_un_tercero_que_no_es_un_nit_se_rechaza(conn):
    asiento.generar(conn, "2026-09")
    a = _uno(conn)
    lineas = [dict(x) for x in a["lineas"]]
    lineas[-1]["tercero_nit"] = "Pepe"
    with pytest.raises(ValueError, match="no es un NIT"):
        asiento.guardar(conn, a["id"], lineas)


# ─────────────────────────── aprobar en bloque ────────────────────────────────

def test_aprobar_en_bloque_solo_los_marcados(conn):
    asiento.generar(conn, "2026-09")
    a = _uno(conn)
    r = asiento.aprobar_lote(conn, "2026-09", ids=[])
    assert r["aprobados"] == 0
    r = asiento.aprobar_lote(conn, "2026-09", ids=[a["id"]])
    assert r["aprobados"] == 1 and _uno(conn)["aprobado"] == 1


def test_en_bloque_se_salta_lo_que_falta_decidir(conn):
    asiento.generar(conn, "2026-09")
    a = _uno(conn)
    lineas = [dict(x) for x in a["lineas"]]
    lineas[0]["cuenta_puc"] = ""
    asiento.guardar(conn, a["id"], lineas)
    r = asiento.aprobar_lote(conn, "2026-09", ids=[a["id"]])
    assert r["aprobados"] == 0 and r["omitidos"] == 1


def test_en_bloque_uno_guardado_aprende_su_regla(conn):
    asiento.generar(conn, "2026-09")
    a = _uno(conn)
    lineas = [dict(x) for x in a["lineas"]]
    lineas[0]["cuenta_puc"] = "51359501"
    asiento.guardar(conn, a["id"], lineas, {"cuenta_puc": "51359501"})
    asiento.aprobar_lote(conn, "2026-09", ids=[a["id"]])
    regla = causacion.proponer(conn, {"tercero_nit": "800100200", "naturaleza": "compra"})
    assert regla.cuenta_puc == "51359501"


def test_aprobar_un_asiento_editado_no_le_borra_la_retencion_a_la_regla(conn):
    concepto = tarifas.listar_conceptos(conn)[0]["id"]
    asiento.generar(conn, "2026-09")
    a = _uno(conn)
    lineas = [dict(x) for x in a["lineas"]]
    lineas[0]["descripcion"] = "Otra descripción"
    asiento.aprobar(conn, a["id"], lineas=lineas)
    regla = causacion.proponer(conn, {"tercero_nit": "800100200", "naturaleza": "compra"})
    assert regla.concepto_retencion_id == concepto


# ───────────── sin XML: la tarifa sale del PDF o se deduce (2026-10-29) ─────────

TABLA_PDF = """
                     [DISCRIMINACIÓN DE IMPUESTOS]
                     DESCRIPCION BASE/CANT IMPUESTO
                         IVA 5%: 40.000,00 2.000,00
                         IVA 0%: 10.000,00      0,00
                        IVA 19%: 10.000,00 1.900,00
"""


def test_la_tabla_de_impuestos_del_pdf_se_lee_y_se_verifica():
    from contab_core import pdf_impuestos as p
    filas = p.desglose_del_texto(TABLA_PDF)
    assert [(str(t), str(b), str(v)) for t, b, v in filas] == [
        ("5", "40000.00", "2000.00"), ("0", "10000.00", "0.00"), ("19", "10000.00", "1900.00")]
    assert p.verificar(filas, subtotal=Decimal("60000"), iva=Decimal("3900"))
    assert not p.verificar(filas, subtotal=Decimal("60000"), iva=Decimal("3950")), \
        "si no da el IVA del reporte, no se usa"
    assert not p.verificar(filas, subtotal=Decimal("50000"), iva=Decimal("3900")), \
        "las bases no pueden pasar del subtotal"


def test_los_montos_se_leen_con_punto_o_coma():
    from contab_core.pdf_impuestos import _monto
    assert _monto("14.487,39") == _monto("14,487.39") == Decimal("14487.39")
    assert _monto("1.500.000") == Decimal("1500000")


def _de_excel(subtotal, iva):
    total = Decimal(subtotal) + Decimal(iva)
    return {"cufe": "c" * 96, "naturaleza": "compra", "subtotal": subtotal,
            "total_impuestos": iva, "total": str(total), "origen": "excel",
            "impuestos": [{"tipo": "01", "base": subtotal, "tarifa": None, "valor": iva}]}


def test_sin_xml_la_factura_mixta_toma_las_tarifas_del_pdf(conn, monkeypatch):
    from contab_core import pdf_impuestos as p
    monkeypatch.setattr(p, "desglose_del_pdf", lambda *a, **k: p.desglose_del_texto(TABLA_PDF))
    doc, aviso = asiento._con_tarifas(conn, {"nit": ""}, _de_excel("60000", "3900"))
    partes = asiento.desglose_de(doc)
    assert [t for t, _b, _v in partes.ivas] == ["19", "5"]
    assert partes.base_gravada == Decimal("50000"), "lo del 0 % queda sin IVA"
    assert "del PDF" in aviso


def test_sin_pdf_una_tarifa_parcial_se_deduce(conn, monkeypatch):
    """La compra de septiembre: 125.700 con IVA de 2.388,30 → base 12.570 al 19 %."""
    from contab_core import pdf_impuestos as p
    monkeypatch.setattr(p, "desglose_del_pdf", lambda *a, **k: None)
    doc, aviso = asiento._con_tarifas(conn, {"nit": ""}, _de_excel("125700", "2388.30"))
    partes = asiento.desglose_de(doc)
    assert partes.ivas == [("19", Decimal("12570.00"), Decimal("2388.30"))]
    assert "19 %" in aviso


def test_si_no_se_puede_saber_la_tarifa_se_dice(conn, monkeypatch):
    from contab_core import pdf_impuestos as p
    monkeypatch.setattr(p, "desglose_del_pdf", lambda *a, **k: None)
    doc, aviso = asiento._con_tarifas(conn, {"nit": ""}, _de_excel("80044.53", "4985.46"))
    assert asiento.desglose_de(doc) is None
    assert "Repártelo en las líneas" in aviso


def test_con_el_19_por_ciento_entero_no_se_busca_nada(conn, monkeypatch):
    from contab_core import pdf_impuestos as p

    def no(*a, **k):
        raise AssertionError("no tenía que abrir el PDF")
    monkeypatch.setattr(p, "desglose_del_pdf", no)
    _doc, aviso = asiento._con_tarifas(conn, {"nit": ""}, _de_excel("100000", "19000"))
    assert aviso is None


def test_los_centavos_del_documento_se_conservan(conn):
    doc = conn.execute("SELECT id FROM documento").fetchone()["id"]
    conn.execute("DELETE FROM documento_impuesto WHERE documento_id = ? AND linea_id IS NULL"
                 " AND es_retencion = 0", (doc,))
    conn.execute("INSERT INTO documento_impuesto(documento_id, tipo, nombre, base, tarifa, valor)"
                 " VALUES (?, '01', 'IVA', '12570', '19.00', '2388.30')", (doc,))
    conn.execute("UPDATE documento SET subtotal = '125700', total_impuestos = '2388.30',"
                 " total = '128088.30' WHERE id = ?", (doc,))
    conn.commit()
    asiento.generar(conn, "2026-09")
    lineas = _uno(conn)["lineas"]
    iva = next(x for x in lineas if x["rol"] == "iva")
    assert Decimal(iva["debito"]) == Decimal("2388.30")
    assert _cuadra(lineas)


# ─────────────────────────── el archivo de Siigo ──────────────────────────────

def test_un_aprobado_que_no_cuadra_no_sale_en_el_archivo(conn, tmp_path):
    asiento.generar(conn, "2026-09")
    a = _uno(conn)
    asiento.aprobar(conn, a["id"])
    # Como uno aprobado en una versión vieja, sin su línea de IVA.
    conn.execute("DELETE FROM asiento_linea WHERE asiento_id = ? AND rol = 'iva'", (a["id"],))
    conn.commit()
    siigo.guardar_config(conn, {"tipo:compra": "11", "ultimo:11": "0"})
    r = siigo.export_siigo(conn, "2026-09", tmp_path)
    assert r["comprobantes"] == 0 and len(r["no_salen"]) == 1
    assert _uno(conn)["siigo_consecutivo"] is None, "no gasta consecutivo"


def test_lo_aprobado_sin_numero_lo_recupera_al_volver_a_armar(conn):
    asiento.generar(conn, "2026-09")
    a = _uno(conn)
    asiento.aprobar(conn, a["id"])
    conn.execute("UPDATE asiento_propuesto SET numero = NULL WHERE id = ?", (a["id"],))
    conn.commit()
    asiento.generar(conn, "2026-09")
    assert _uno(conn)["numero"]
