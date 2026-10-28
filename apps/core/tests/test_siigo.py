"""El archivo de importación de comprobantes de Siigo (D-128).

Lo que estas pruebas vigilan: que salga solo lo aprobado, con las 27 columnas
del modelo; que cada comprobante lleve su tipo, su consecutivo y el número de
la factura en la línea de cartera; que las líneas de IVA y retención lleven
código y base; que exportar dos veces no cambie un consecutivo; y que lo que
falte se diga con nombre en vez de inventarse.
"""
from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook

from contab_core import (
    asiento,
    causacion,
    documentos,
    perfil_fiscal,
    siigo,
    storage,
    tarifas,
    xml_ubl,
)

MUESTRAS = Path(__file__).resolve().parents[3] / "docs" / "samples" / "xml"
CUFE = "1" * 92 + "aaaa"


@pytest.fixture
def conn(tmp_path: Path):
    db = tmp_path / "e.db"
    storage.init_empresa_db(db, nit="900181831", dv="4", razon_social="REAL")
    c = storage.connect(db)
    storage.ensure_schema(c)
    perfil_fiscal.guardar_perfil(c, {
        "vigencia_desde": "2026-01-01", "regimen": "ordinario", "responsable_iva": 1,
        "periodicidad_iva": "cuatrimestral", "agente_retefuente": 1, "agente_reteiva": 1})
    tarifas.guardar_uvt(c, 2026, "52374")
    tarifas.guardar_concepto(c, {
        "concepto": "Compras.", "aplica_a": "ambos", "sobre": "base", "tarifa": "2.5",
        "base_minima_uvt": "10", "cuenta_puc": "23654001", "codigo_siigo": "14",
        "vigencia_desde": "2026-01-01"})
    lectura = xml_ubl.leer(MUESTRAS / "factura-con-iva.xml")
    documentos.guardar(c, lectura.documento, MUESTRAS / "factura-con-iva.xml")
    c.execute("INSERT INTO excel_load(id, source_path, source_name, total_rows, periodo)"
              " VALUES (1,'x.xlsx','x.xlsx',1,'2026-09')")
    c.execute("INSERT INTO excel_row(load_id,row_index,cufe_cude,data_json,periodo)"
              " VALUES (1,0,?,'{\"Grupo\":\"Recibido\"}','2026-09')", (CUFE,))
    c.execute("UPDATE empresa_info SET cuenta_proveedores='22050501', cuenta_iva_descontable='24081501'")
    c.commit()
    yield c
    c.close()


def _filas(path):
    ws = load_workbook(path).active
    return [[cell.value for cell in row] for row in ws.iter_rows()]


def _aprobar_la_compra(conn):
    asiento.generar(conn, "2026-09")
    a = asiento.listar(conn, "2026-09")[0]
    concepto = tarifas.listar_conceptos(conn)[0]["id"]
    asiento.aprobar(conn, a["id"], cuenta_puc="51953001", concepto_retencion_id=concepto)
    return a["id"]


def test_partir_el_numero_de_la_factura():
    assert siigo.partir_numero("FE1-123456") == ("FE1", "123456")
    assert siigo.partir_numero("DSE 242") == ("DSE", "242")
    assert siigo.partir_numero("FELE123") == ("FELE", "123")
    assert siigo.partir_numero("845") == ("", "845")
    assert siigo.partir_numero(None) == ("", "")


def test_sin_configurar_sale_el_archivo_y_se_dice_que_falta(conn, tmp_path):
    _aprobar_la_compra(conn)
    r = siigo.export_siigo(conn, "2026-09", tmp_path)
    assert Path(r["path"]).name == "siigo-2026-09.xlsx"
    assert "tipo:compra" in r["faltan"] and "iva:19" in r["faltan"]
    assert r["comprobantes"] == 1 and r["cuadra"] is True


def test_el_comprobante_sale_como_lo_espera_siigo(conn, tmp_path):
    siigo.guardar_config(conn, {"tipo:compra": "11", "ultimo:11": "849", "iva:19": "02"})
    _aprobar_la_compra(conn)
    r = siigo.export_siigo(conn, "2026-09", tmp_path)
    assert r["faltan"] == []
    filas = _filas(r["path"])
    assert filas[0] == siigo.COLUMNAS
    cuerpo = filas[1:]
    assert len(cuerpo) == r["lineas"] == 4
    assert {f[0] for f in cuerpo} == {11}, "tipo de comprobante"
    assert {f[1] for f in cuerpo} == {850}, "sigue al último consecutivo"
    assert {f[6] for f in cuerpo} == {800100200}, "el tercero en todas las líneas"
    por_cuenta = {f[5]: f for f in cuerpo}
    iva = por_cuenta[24081501]
    assert iva[16] == "02" and iva[24] == 1500000, "código de impuesto y base en la línea de IVA"
    ret = por_cuenta[23654001]
    assert ret[16] == "14" and ret[24] == 1500000
    cartera = por_cuenta[22050501]
    assert cartera[12] and cartera[13] and cartera[14] == 1, "prefijo, consecutivo y cuota en la cartera"
    assert (cartera[15].date() - cartera[2].date()).days == 30
    assert por_cuenta[51953001][12] in ("", None), "el gasto no lleva prefijo"
    assert sum(f[21] or 0 for f in cuerpo) == sum(f[22] or 0 for f in cuerpo)


def test_exportar_dos_veces_no_cambia_el_consecutivo(conn, tmp_path):
    siigo.guardar_config(conn, {"tipo:compra": "11", "ultimo:11": "849", "iva:19": "02"})
    aid = _aprobar_la_compra(conn)
    siigo.export_siigo(conn, "2026-09", tmp_path)
    primero = asiento.listar_uno(conn, aid)["siigo_consecutivo"]
    siigo.export_siigo(conn, "2026-09", tmp_path)
    assert asiento.listar_uno(conn, aid)["siigo_consecutivo"] == primero == 850
    assert siigo.config(conn)["ultimo:11"] == "850", "el último queda anotado para el mes que viene"


def test_lo_no_aprobado_no_sale(conn, tmp_path):
    siigo.guardar_config(conn, {"tipo:compra": "11"})
    asiento.generar(conn, "2026-09")
    r = siigo.export_siigo(conn, "2026-09", tmp_path)
    assert r["comprobantes"] == 0 and r["lineas"] == 0


def test_la_nota_credito_va_en_positivo_con_las_columnas_al_reves(conn, tmp_path):
    siigo.guardar_config(conn, {"tipo:venta": "12", "tipo:nc_venta": "17", "iva:19": "02"})
    conn.execute("UPDATE empresa_info SET cuenta_clientes='13050501', cuenta_iva_generado='24080504'")
    conn.execute(
        "INSERT INTO excel_row(load_id,row_index,cufe_cude,tipo_doc,data_json,periodo) VALUES (1,1,?,?,?,'2026-09')",
        ("n" * 96, "Nota de crédito electrónica",
         '{"Grupo":"Emitido","NIT Emisor":"900181831","NIT Receptor":"901111111","Nombre Receptor":"CLIENTE",'
         '"Prefijo":"NCRS","Folio":"136","Fecha Emisión":"2026-09-14","Total":"119000","IVA":"19000"}'))
    conn.commit()
    causacion.guardar(conn, {"tercero_nit": "901111111", "cuenta_puc": "41750502", "naturaleza": "venta"})
    asiento.generar(conn, "2026-09")
    nc = next(a for a in asiento.listar(conn, "2026-09") if a["cufe"] == "n" * 96)
    asiento.aprobar(conn, nc["id"], cuenta_puc="41750502")
    r = siigo.export_siigo(conn, "2026-09", tmp_path)
    filas = [f for f in _filas(r["path"])[1:] if f[0] == 17]
    assert len(filas) == 3
    assert all((f[21] or 0) >= 0 and (f[22] or 0) >= 0 for f in filas), "nada en negativo"
    cartera = next(f for f in filas if f[5] == 13050501)
    assert cartera[22] == 119000 and cartera[12] == "NCRS" and cartera[13] == 136
    assert cartera[19] == "NCRS-136"
    assert cartera[2].date() == date(2026, 9, 14)
    assert cartera[15].date() == date(2026, 10, 14), "vence a los 30 días"


# ───────────────── el catálogo de Siigo, desde el Excel de cada empresa ─────────

def _excel_catalogo(ruta: Path):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Impuestos"
    ws.append(["Código", "Nombre", "Tipo", "Tarifa (%)", "Cuenta"])
    ws.append([2, "IVA 19%", "IVA", 19, 24081501])
    ws.append(["01", "IVA 5%", "IVA", "5", "24081001"])
    ws.append([14, "Retención compras 3,5%", "Retefuente", 3.5, 23653003])
    ws.append(["09", "Retención compras 2,5%", "Retefuente", "2,5", ""])
    ws.append(["07", "ReteIVA", "ReteIVA", 15, 13551703])
    ws.append(["100", "Autorretención", "Autorretención", 0.8, 23657501])
    ws2 = wb.create_sheet("Comprobantes")
    ws2.append(["Clase", "Tipo en Siigo", "Último consecutivo"])
    ws2.append(["Factura de compra", 11, 874])
    ws2.append(["Nota crédito de compra", 15, 22])
    ws2.append(["Factura de venta", 12, 439])
    ws2.append(["Nota crédito de venta", 17, 140])
    ws2.append(["Documento soporte", 19, 262])
    ws2.append(["Depreciación", 2, 9])
    wb.save(ruta)


def test_el_catalogo_se_lee_y_se_cruza_con_los_conceptos(conn, tmp_path):
    ruta = tmp_path / "Siigo REAL.xlsx"
    _excel_catalogo(ruta)
    cargado = siigo.cargar_catalogo(ruta)
    assert [i["codigo"] for i in cargado.impuestos] == ["02", "01", "14", "09", "07", "100"]
    assert cargado.impuestos[0]["tarifa"] == "19" and cargado.impuestos[3]["tarifa"] == "2.5"
    assert any("Depreciacion".lower() in a.lower() or "depreciación" in a.lower() for a in cargado.avisos)

    vista = siigo.aplicar_catalogo(conn, cargado, dry_run=True)
    assert vista["config"]["iva:19"] == "02" and vista["config"]["iva:5"] == "01"
    assert vista["config"]["tipo:compra"] == "11" and vista["config"]["ultimo:11"] == "874"
    assert vista["config"]["tipo:doc_soporte"] == "19"
    # el concepto de la prueba es «Compras.» al 2,5 % con cuenta 23654001: cruza por tarifa
    assert [c["codigo"] for c in vista["conceptos_cruzados"]] == ["09"]
    assert {x["codigo"] for x in vista["sin_cruce"]} == {"14", "07", "100"}
    assert siigo.config(conn) == {}, "el ensayo no escribe"

    siigo.aplicar_catalogo(conn, cargado)
    assert siigo.config(conn)["tipo:nc_venta"] == "17"
    assert tarifas.listar_conceptos(conn)[0]["codigo_siigo"] == "09"


def test_con_el_catalogo_cargado_el_archivo_sale_completo(conn, tmp_path):
    ruta = tmp_path / "Siigo REAL.xlsx"
    _excel_catalogo(ruta)
    siigo.aplicar_catalogo(conn, siigo.cargar_catalogo(ruta))
    _aprobar_la_compra(conn)
    r = siigo.export_siigo(conn, "2026-09", tmp_path)
    assert r["faltan"] == []
    filas = _filas(r["path"])[1:]
    assert {f[1] for f in filas} == {875}, "sigue al 874 del catálogo"
    ret = next(f for f in filas if f[5] == 23654001)
    assert ret[16] == "09"
