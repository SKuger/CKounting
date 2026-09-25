"""Tests del PUC: parseo del dump, upsert (fusionar), CRUD manual, bloqueo por NIT."""
from pathlib import Path

import pytest
from openpyxl import Workbook

from contab_core import puc, storage
from contab_core import sidecar as S

HEADER = ["Código", "Nombre", "Categoría", "Clase", "Relación con",
          "Maneja vencimientos", "Diferencia fiscal", "Activo", "Nivel agrupación"]


def _make_dump(path: Path, nit: str = "900123456-7", rows=None) -> Path:
    """Escribe un Excel con el layout real: preámbulo + encabezado + filas + pie."""
    if rows is None:
        rows = [
            ("1", "Activo", None, None, None, None, None, None, None),
            ("11", "Efectivo y equivalentes", None, None, None, None, None, None, None),
            ("1105", "Caja", None, None, None, None, None, None, None),
            ("110505", "Caja general", None, None, None, None, None, None, None),
            ("11050501", "Caja general", "Caja - Bancos", "Activo", "Sin asignar",
             "No maneja vencimiento", "No", "Sí", "Transaccional"),
            ("2105", "Bancos nacionales", None, None, None, None, None, "No", None),  # Activo=No
        ]
    wb = Workbook()
    ws = wb.active
    ws.append(["Cuentas contables"])
    ws.append(["EMPRESA DE PRUEBA S.A.S."])
    ws.append([nit])
    ws.append([])
    ws.append(HEADER)
    for r in rows:
        ws.append(list(r))
    ws.append([])
    ws.append(["Procesado en: Agosto 09 2026 11:04"])
    wb.save(path)
    return path


@pytest.mark.parametrize("codigo,nivel", [("1", 1), ("11", 2), ("1105", 3), ("110505", 4), ("11050501", 5), ("1105050101", 6)])
def test_nivel(codigo, nivel):
    assert puc.puc_nivel(codigo) == nivel


def test_load_puc_parsea_preambulo_y_filas(tmp_path):
    loaded = puc.load_puc(_make_dump(tmp_path / "puc.xlsx"))
    assert loaded.empresa_nit == "900123456"          # NIT sin DV
    assert "EMPRESA DE PRUEBA" in loaded.empresa_nombre
    assert len(loaded.accounts) == 6                   # el pie 'Procesado en:' se descarta
    codigos = [a["codigo"] for a in loaded.accounts]
    assert "11050501" in codigos
    hoja = next(a for a in loaded.accounts if a["codigo"] == "11050501")
    assert hoja["nivel"] == 5 and hoja["nivel_agrupacion"] == "Transaccional" and hoja["activo"] == 1
    assert next(a for a in loaded.accounts if a["codigo"] == "2105")["activo"] == 0  # Activo=No


def _empresa(tmp_path, nit="900123456", dv="7"):
    db = tmp_path / "empresa.db"
    storage.init_empresa_db(db, nit=nit, dv=dv, razon_social="EMPRESA DE PRUEBA")
    return str(db)


def test_persist_merge_upsert(tmp_path):
    db = _empresa(tmp_path)
    src = _make_dump(tmp_path / "puc.xlsx")
    # primera importación → todo nuevo
    st = S._h_puc_import({"db_path": db, "source_path": str(src)})
    assert st["inserted"] == 6 and st["updated"] == 0
    # cuenta manual que NO viene en el dump
    S._h_puc_save_account({"db_path": db, "account": {"codigo": "99990001", "nombre": "Manual"}})
    # re-importar el mismo dump → actualiza, no duplica, y conserva la manual
    st2 = S._h_puc_import({"db_path": db, "source_path": str(src)})
    assert st2["inserted"] == 0 and st2["updated"] == 6
    data = S._h_puc_list({"db_path": db})
    assert data["counts"]["total"] == 7          # 6 del dump + 1 manual
    assert data["counts"]["manuales"] == 1


def test_dry_run_no_escribe(tmp_path):
    db = _empresa(tmp_path)
    src = _make_dump(tmp_path / "puc.xlsx")
    st = S._h_puc_import({"db_path": db, "source_path": str(src), "dry_run": True})
    assert st["inserted"] == 6 and st["dry_run"] is True
    assert S._h_puc_list({"db_path": db})["counts"]["total"] == 0  # no se escribió nada


def test_bloqueo_por_nit(tmp_path):
    db = _empresa(tmp_path, nit="811000000")     # empresa con OTRO nit
    src = _make_dump(tmp_path / "puc.xlsx", nit="900123456-7")
    with pytest.raises(ValueError, match="no coincide"):
        S._h_puc_import({"db_path": db, "source_path": str(src)})
    assert S._h_puc_list({"db_path": db})["counts"]["total"] == 0


def test_crud_manual_y_archivar(tmp_path):
    db = _empresa(tmp_path)
    conn = storage.connect(db)
    storage.ensure_schema(conn)
    acc = puc.save_account(conn, {"codigo": "13050501", "nombre": "Clientes X"})
    assert acc["fuente"] == "manual" and acc["nivel"] == 5
    # editar
    acc2 = puc.save_account(conn, {"codigo": "13050501", "nombre": "Clientes nacionales"}, account_id=acc["id"])
    assert acc2["nombre"] == "Clientes nacionales"
    # código duplicado → error
    with pytest.raises(puc.PucParseError, match="Ya existe"):
        puc.save_account(conn, {"codigo": "13050501", "nombre": "Otra"})
    # archivar
    puc.set_archived(conn, acc["id"], True)
    assert len(puc.list_accounts(conn, include_archived=False)) == 0
    assert len(puc.list_accounts(conn, include_archived=True)) == 1
    # eliminar
    puc.delete_account(conn, acc["id"])
    assert puc.counts(conn)["total"] == 0
    conn.close()


@pytest.mark.parametrize("cell,nit", [
    ("900181831-4", "900181831"),
    ("NIT: 900.181.831-4", "900181831"),
    ("Nit 900181831", "900181831"),
    ("20260809", ""),                 # fecha: no se confunde con NIT
    ("Cuentas contables", ""),
])
def test_extract_nit(cell, nit):
    assert puc._extract_nit(cell) == nit


def test_bloqueo_con_nit_etiquetado(tmp_path):
    """El bloqueo por NIT (D-003/D-070) también funciona si el Excel trae el NIT
    con etiqueta/guión, no solo como celda de dígitos puros."""
    db = _empresa(tmp_path, nit="811000000")
    src = _make_dump(tmp_path / "puc.xlsx", nit="NIT: 900123456-7")
    with pytest.raises(ValueError, match="no coincide"):
        S._h_puc_import({"db_path": db, "source_path": str(src)})


def test_header_en_columna_no_cero(tmp_path):
    """Un export con una columna vacía a la izquierda igual se parsea (el 'Código'
    puede estar en cualquier columna)."""
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.append([None, "Cuentas contables"])
    ws.append([None, "EMPRESA X"])
    ws.append([None, "900123456-7"])
    ws.append([None, *HEADER])                       # encabezado corrido una columna
    ws.append([None, "1", "Activo", None, None, None, None, None, None])
    ws.append([None, "1105", "Caja", None, None, None, None, None, None])
    path = tmp_path / "corrido.xlsx"
    wb.save(path)
    loaded = puc.load_puc(path)
    assert loaded.empresa_nit == "900123456"
    assert [a["codigo"] for a in loaded.accounts] == ["1", "1105"]


def test_export_roundtrip(tmp_path):
    db = _empresa(tmp_path)
    S._h_puc_import({"db_path": db, "source_path": str(_make_dump(tmp_path / "puc.xlsx"))})
    out = tmp_path / "export.xlsx"
    res = S._h_puc_export({"db_path": db, "output_path": str(out)})
    assert res["rows_written"] == 6 and out.exists()
    from openpyxl import load_workbook
    wb = load_workbook(out, read_only=True)
    data = list(wb.active.iter_rows(values_only=True))
    wb.close()
    assert data[0] == tuple(HEADER)
