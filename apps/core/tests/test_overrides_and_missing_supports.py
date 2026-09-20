"""D-046 (overrides) + D-047 (faltantes de soporte report)."""
from pathlib import Path

import pytest
from openpyxl import load_workbook

from contab_core import excel, storage
from contab_core.export import MISSING_SUPPORTS_COLUMNS
from contab_core.sidecar import HANDLERS


@pytest.fixture
def empresa_with_load(tmp_path: Path, sample_excel: Path) -> dict:
    db_path = tmp_path / "empresa.db"
    storage.init_empresa_db(db_path, nit="900181831", dv="4", razon_social="REAL")
    loaded = excel.load_excel(sample_excel)
    conn = storage.connect(db_path)
    try:
        load_id = excel.persist_load(conn, str(sample_excel), loaded)
    finally:
        conn.close()
    return {"db_path": db_path, "load_id": load_id, "rows": loaded.rows}


# ---------- D-046: overrides ----------


def test_toggle_override_persists(empresa_with_load):
    db_path = empresa_with_load["db_path"]
    load_id = empresa_with_load["load_id"]

    res = HANDLERS["toggle_validation_override"]({
        "db_path": str(db_path), "load_id": load_id,
        "row_index": 3, "code": "V3", "active": True,
    })
    assert 3 in res["v3_overrides"]

    # Idempotente.
    again = HANDLERS["toggle_validation_override"]({
        "db_path": str(db_path), "load_id": load_id,
        "row_index": 3, "code": "V3", "active": True,
    })
    assert again["v3_overrides"].count(3) == 1

    # Off → desaparece.
    off = HANDLERS["toggle_validation_override"]({
        "db_path": str(db_path), "load_id": load_id,
        "row_index": 3, "code": "V3", "active": False,
    })
    assert 3 not in off["v3_overrides"]


def test_get_load_rows_returns_overrides(empresa_with_load):
    db_path = empresa_with_load["db_path"]
    load_id = empresa_with_load["load_id"]
    HANDLERS["toggle_validation_override"]({
        "db_path": str(db_path), "load_id": load_id,
        "row_index": 5, "code": "V3", "active": True,
    })
    res = HANDLERS["get_load_rows"]({
        "db_path": str(db_path), "load_id": load_id, "offset": 0, "limit": 100,
    })
    assert res["v3_overrides"] == [5]


def test_run_validations_respects_override(empresa_with_load):
    db_path = empresa_with_load["db_path"]
    load_id = empresa_with_load["load_id"]

    # Run with no override → BEC-509558651 fails V3 (off-rate).
    before = HANDLERS["run_validations"]({
        "db_path": str(db_path), "load_id": load_id,
    })
    bec_idx = next(
        i for i, r in enumerate(empresa_with_load["rows"])
        if r["# de factura"] == "BEC-509558651"
    )
    bec_codes_before = next(
        (f["codes"] for f in before["failures"] if f["row_index"] == bec_idx),
        [],
    )
    assert "V3" in bec_codes_before

    # Override → BEC-509558651 ya no falla V3.
    HANDLERS["toggle_validation_override"]({
        "db_path": str(db_path), "load_id": load_id,
        "row_index": bec_idx, "code": "V3", "active": True,
    })
    after = HANDLERS["run_validations"]({
        "db_path": str(db_path), "load_id": load_id,
    })
    bec_codes_after = next(
        (f["codes"] for f in after["failures"] if f["row_index"] == bec_idx),
        [],
    )
    assert "V3" not in bec_codes_after


# ---------- D-047: faltantes de soporte ----------


def test_missing_supports_report_has_5_columns(empresa_with_load, tmp_path: Path):
    db_path = empresa_with_load["db_path"]
    load_id = empresa_with_load["load_id"]
    out = tmp_path / "faltantes.xlsx"

    # Configurar carpeta de PDFs vacía → todas las filas fallan V1.
    pdfs_folder = tmp_path / "pdfs_vacia"
    pdfs_folder.mkdir()
    storage.set_periodo_folder(db_path, "2026-02", str(pdfs_folder))  # 0.1.4: soportes por mes

    res = HANDLERS["export_missing_supports"]({
        "db_path": str(db_path), "load_id": load_id,
        "output_path": str(out),
    })
    assert res["rows_written"] > 0
    assert out.exists()

    wb = load_workbook(out, read_only=False, data_only=True)
    sheet = wb.active
    headers = [c.value for c in sheet[1]]
    assert headers == MISSING_SUPPORTS_COLUMNS
    assert len(headers) == 5
    wb.close()


def test_missing_supports_export_inverts_nc_totals(empresa_with_load, tmp_path: Path):
    db_path = empresa_with_load["db_path"]
    load_id = empresa_with_load["load_id"]
    out = tmp_path / "faltantes_nc.xlsx"
    pdfs_folder = tmp_path / "pdfs_vacia"
    pdfs_folder.mkdir()
    storage.set_periodo_folder(db_path, "2026-02", str(pdfs_folder))  # 0.1.4: soportes por mes

    HANDLERS["export_missing_supports"]({
        "db_path": str(db_path), "load_id": load_id,
        "output_path": str(out),
    })
    wb = load_workbook(out, read_only=False, data_only=True)
    sheet = wb.active
    headers = [c.value for c in sheet[1]]
    factura_idx = headers.index("# de factura")
    total_idx = headers.index("Total")
    saw_nc_negative = False
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if row[factura_idx] in ("TCNC-1952292", "NCLA-1017"):
            assert row[total_idx] <= 0, f"NC {row[factura_idx]} should have negative Total"
            saw_nc_negative = True
    wb.close()
    assert saw_nc_negative, "should have at least one NC in the report"


def test_missing_supports_errors_when_no_pdfs_folder(empresa_with_load, tmp_path: Path):
    db_path = empresa_with_load["db_path"]
    load_id = empresa_with_load["load_id"]
    out = tmp_path / "no_folder.xlsx"
    # No pdfs_folder set → must error out (V1 not active).
    with pytest.raises(Exception) as exc_info:
        HANDLERS["export_missing_supports"]({
            "db_path": str(db_path), "load_id": load_id,
            "output_path": str(out),
        })
    assert "carpeta" in str(exc_info.value).lower()
