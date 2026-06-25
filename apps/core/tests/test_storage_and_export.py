import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from contab_core import excel, export, storage, validations


@pytest.fixture
def empresa_db(tmp_path: Path) -> Path:
    db_path = tmp_path / "test_empresa.db"
    storage.init_empresa_db(db_path, nit="900181831", dv="4", razon_social="REAL Seguros")
    return db_path


def test_empresa_info_persisted(empresa_db):
    info = storage.read_empresa_info(empresa_db)
    assert info["nit"] == "900181831"
    assert info["dv"] == "4"
    assert info["razon_social"] == "REAL Seguros"


def test_persist_and_fetch_load(empresa_db, sample_excel):
    loaded = excel.load_excel(sample_excel)
    conn = storage.connect(empresa_db)
    try:
        load_id = excel.persist_load(conn, str(sample_excel), loaded)
        fetched = excel.fetch_load_rows(conn, load_id)
    finally:
        conn.close()
    assert len(fetched) == len(loaded.rows)
    # data_json round-trips with derived columns intact
    for original, retrieved in zip(loaded.rows, fetched):
        assert retrieved["# de factura"] == original["# de factura"]
        assert abs(retrieved["Valor Base"] - original["Valor Base"]) < 0.01


def test_export_nc_inverts_signs(tmp_path: Path, sample_excel):
    loaded = excel.load_excel(sample_excel)
    tagged = validations.evaluate_rows(loaded.rows, run_v3=True, indexed_cufes=set())
    reasons = validations.failures_to_strings(tagged)
    output = tmp_path / "no_conformes.xlsx"
    result = export.export_no_conformes(loaded.rows, reasons, output)
    assert result["rows_written"] > 0
    assert output.exists()

    wb = load_workbook(output, read_only=False, data_only=True)
    sheet = wb.active
    headers = [c.value for c in sheet[1]]
    assert "razon_falla" in headers
    assert "# de factura" in headers
    assert "Valor Base" in headers
    assert "Base IVA" in headers
    assert "Diferencia" in headers
    # 32 DIAN + 4 derived (#factura, Valor Base, Base IVA, Diferencia) + razon_falla = 37.
    assert len(headers) == 37

    # Find a NC row in the output and assert its IVA and Total are negative.
    tipo_idx = headers.index("Tipo de documento")
    iva_idx = headers.index("IVA")
    total_idx = headers.index("Total")
    for excel_row in sheet.iter_rows(min_row=2, values_only=True):
        tipo = excel_row[tipo_idx]
        if isinstance(tipo, str) and tipo.lower().startswith("nota de cr"):
            assert excel_row[iva_idx] <= 0
            assert excel_row[total_idx] <= 0
            break
    wb.close()


def test_export_uses_semicolon_separator(tmp_path: Path, sample_excel):
    loaded = excel.load_excel(sample_excel)
    # Force multiple reasons by enabling both checks; many sample rows will fail V3
    # AND get the V1 stub at the same time.
    tagged = validations.evaluate_rows(loaded.rows, run_v3=True, indexed_cufes=set())
    reasons = validations.failures_to_strings(tagged)
    output = tmp_path / "multi.xlsx"
    export.export_no_conformes(loaded.rows, reasons, output)

    wb = load_workbook(output, read_only=False, data_only=True)
    sheet = wb.active
    headers = [c.value for c in sheet[1]]
    razon_idx = headers.index("razon_falla")
    found_semicolon = False
    for excel_row in sheet.iter_rows(min_row=2, values_only=True):
        value = excel_row[razon_idx]
        if value and "; " in value:
            found_semicolon = True
            break
    wb.close()
    assert found_semicolon, "Expected at least one row with multiple reasons joined by '; '"
