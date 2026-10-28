"""La nómina como comprobante de Siigo (D-134).

El caso real que estas pruebas reproducen en pequeño: el Excel de nómina del
contador (hojas 1Q/2Q, cabecera en tres filas) y el 992 que él ya subió a Siigo
para esa quincena. Contab aprende la plantilla comparando los dos, y con ella
arma un 992 igual al suyo. Sobre el septiembre real de REAL, la segunda
quincena sale idéntica en sus 80 pares cuenta-tercero.
"""
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from contab_core import nomina, siigo, storage

EMPLEADOS = [
    # id, nombre, salario quincena, aux transporte, salud, pensión, fsp, retención, otras ded, rodamiento, otros pagos
    (98562771, "EMPLEADO UNO", 4300000, 0, 172000, 172000, 43000, 175000, 0, 0, 1000000),
    (43989893, "EMPLEADA DOS", 1850000, 124547.49999999999, 74000, 74000, 0, 0, 50000, 600000, 0),
]
EPS, FONDO, CAJA, ARL, DIAN = "800088702", "900336004", "890900842", "890903790", "800197268"


def _neto(e):
    """Como en su Excel: con el medio peso del auxilio adentro."""
    _id, _n, sal, aux, salud, pen, fsp, ret, otras, rod, otros = e
    return sal + aux + rod + otros - salud - pen - fsp - ret - otras


def _al_peso(x):
    return int(Decimal(str(round(x, 6))).quantize(Decimal("1"), rounding="ROUND_HALF_UP"))


def _excel_nomina(ruta: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "1Q"
    for _ in range(7):
        ws.append([])
    cab = [None, "Id", "Nombre del\nempleado", "Devengado"] + [None] * 23 + ["Deducciones"] + [None] * 7 + ["Otros pagos", None, None, "Neto  pagado", "Observaciones"]
    sub = [None, None, None, "Salario \nbásico", "Días / Horas\nliquidados", "Salario \ndevengado", "Vacaciones", None,
           "Comisones", "Incapacidades", None] + [None] * 14 + ["Auxilio de\ntransporte", "Total\ndevengado", "Salud",
           "Pensión", "Fondo de solidaridad\npensional", "Retención en \nla fuente", "Cartera", "Prestamos",
           "Otras deducciones", "Total Deducciones", "Auxilio de datos", "Auxilio de rodamiento", "Otros pagos no salariales"]
    sub2 = [None] * 6 + ["Días Disfrutados", "Vacaciones", None, "Días incapacidad", "Vr. Incapacidad"]
    ws.append(cab)
    ws.append(sub)
    ws.append(sub2)
    for e in EMPLEADOS:
        _id, nombre, sal, aux, salud, pen, fsp, ret, otras, rod, otros = e
        fila = [None, _id, nombre, sal * 2, 15, sal] + [0] * 19 + [aux, sal + aux, salud, pen, fsp, ret, 0, 0, otras,
                                                                  salud + pen + fsp + ret + otras, None, rod, otros, _neto(e), ""]
        ws.append(fila)
    wb.save(ruta)


def _carga_992(ruta: Path):
    """El 992 de esa quincena tal como lo arma el contador."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Hoja historicos"
    ws.append(siigo.COLUMNAS)

    def linea(cuenta, tercero, deb=0, cre=0):
        fila = [None] * len(siigo.COLUMNAS)
        fila[0], fila[1], fila[2], fila[5], fila[6], fila[21], fila[22] = 992, 65, "2026-09-15", cuenta, tercero, deb, cre
        ws.append(fila)

    p = lambda x: int(Decimal(x).quantize(Decimal("1"), rounding="ROUND_HALF_UP"))  # noqa: E731
    for e in EMPLEADOS:
        _id, _n, sal, aux, salud, pen, fsp, ret, otras, rod, otros = e
        ces = p(Decimal(sal) * Decimal("0.0833"))
        linea(51050601, _id, deb=sal)
        if aux:
            linea(51052701, _id, deb=_al_peso(aux))
        if rod:
            linea(51054503, _id, deb=rod)
        if otros:
            linea(51054503, _id, deb=otros)
        for d, c, v in ((51053001, 2510100101, ces), (51053301, 2510100201, p(ces * Decimal("0.12"))),
                        (51053601, 2510100401, ces), (51053901, 2510100301, p(Decimal(sal) * Decimal("0.0417")))):
            linea(d, _id, deb=v)
            linea(c, _id, cre=v)
        er, caja, arl = p(Decimal(sal) * Decimal("0.12")), p(Decimal(sal) * Decimal("0.04")), p(Decimal(sal) * Decimal("0.00522"))
        linea(51057001, _id, deb=er)
        linea(51057201, _id, deb=caja)
        linea(51056801, _id, deb=arl)
        linea(23700501, EPS, cre=salud)
        linea(23704505, FONDO, cre=pen)
        linea(23704505, FONDO, cre=er)
        if fsp:
            linea(23704505, FONDO, cre=fsp)
        linea(23701001, CAJA, cre=caja)
        linea(23700601, ARL, cre=arl)
        if ret:
            linea(23650501, _id, cre=ret)
        if otras:
            linea(13309501, _id, cre=otras)
        linea(25050501, _id, cre=_al_peso(_neto(e)))
    wb.save(ruta)


@pytest.fixture
def conn(tmp_path: Path):
    db = tmp_path / "e.db"
    storage.init_empresa_db(db, nit="900181831", dv="4", razon_social="REAL")
    c = storage.connect(db)
    storage.ensure_schema(c)
    siigo.guardar_config(c, {"tipo:nomina": "992", "ultimo:992": "66"})
    yield c
    c.close()


@pytest.fixture
def archivos(tmp_path: Path):
    n, k = tmp_path / "NOMINA.xlsx", tmp_path / "CARGA.xlsx"
    _excel_nomina(n)
    _carga_992(k)
    return n, k


def test_el_excel_de_nomina_se_lee_por_concepto(archivos):
    q = nomina.leer_nomina(archivos[0])["1Q"]
    assert [e.id for e in q.empleados] == ["98562771", "43989893"]
    dos = q.empleados[1]
    assert dos.v("salario") == 1850000 and dos.v("aux_transporte") == 124548, "medio peso, al peso hacia arriba"
    assert dos.v("rodamiento") == 600000 and dos.v("otras_deducciones") == 50000
    assert q.columnas_sin_leer == []


def test_la_plantilla_se_aprende_del_992_que_ya_esta_en_siigo(conn, archivos):
    r = nomina.aprender(conn, *archivos)
    assert r["faltan"] == []
    cuentas = r["cuentas"]
    assert cuentas["nom:debito:salario"] == "51050601"
    assert (cuentas["nom:debito:cesantias"], cuentas["nom:debito:prima"]) == ("51053001", "51053601")
    assert (cuentas["nom:credito:cesantias"], cuentas["nom:credito:prima"]) == ("2510100101", "2510100401")
    assert cuentas["nom:credito:otras_deducciones"] == "13309501"
    assert cuentas["nom:credito:neto"] == "25050501"
    assert f"23700501|{EPS}" in r["entidades"] and f"23704505|{FONDO}" in r["entidades"]


def test_el_992_generado_es_igual_al_suyo_por_cuenta_y_tercero(conn, archivos):
    nomina.aprender(conn, *archivos)
    g = nomina.generar(conn, "2026-09", archivos[0])
    [q] = g["quincenas"]
    assert q["cuadra"] and q["avisos"] == []
    nuestro: dict = {}
    for ln in nomina.listar(conn, "2026-09")[0]["lineas"]:
        k = (ln["cuenta"], ln["tercero"])
        nuestro[k] = nuestro.get(k, 0) + Decimal(ln["debito"]) - Decimal(ln["credito"])
    suyo: dict = {}
    for f in siigo._filas_siigo(archivos[1]):
        k = (str(f["Código cuenta contable"]), str(f["Identificación tercero"]))
        suyo[k] = suyo.get(k, 0) + Decimal(str(f.get("Débito") or 0)) - Decimal(str(f.get("Crédito") or 0))
    assert nuestro == suyo


def test_sin_plantilla_no_se_aprueba(conn, archivos):
    nomina.generar(conn, "2026-09", archivos[0])
    c = nomina.listar(conn, "2026-09")[0]
    assert any("Falta" in a for a in c["avisos"])
    with pytest.raises(ValueError, match="sin cuenta"):
        nomina.aprobar(conn, c["id"])


def test_la_nomina_aprobada_sale_en_el_archivo_de_siigo(conn, archivos, tmp_path):
    nomina.aprender(conn, *archivos)
    nomina.generar(conn, "2026-09", archivos[0])
    c = nomina.listar(conn, "2026-09")[0]
    nomina.aprobar(conn, c["id"])
    r = siigo.export_siigo(conn, "2026-09", tmp_path / "out")
    assert r["cuadra"] and r["comprobantes"] == 1
    filas = [[cell.value for cell in row] for row in load_workbook(r["path"]).active.iter_rows()][1:]
    assert {f[0] for f in filas} == {992} and {f[1] for f in filas} == {67}, "sigue al 66"
    assert {f[6] for f in filas} >= {98562771, int(EPS), int(FONDO)}, "cada línea con su tercero"
    # exportar otra vez no cambia el consecutivo ni rehace lo aprobado
    nomina.generar(conn, "2026-09", archivos[0])
    siigo.export_siigo(conn, "2026-09", tmp_path / "out")
    assert nomina.listar(conn, "2026-09")[0]["siigo_consecutivo"] == 67
