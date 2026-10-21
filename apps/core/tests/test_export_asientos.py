"""Exportar los asientos aprobados a Excel.

Es lo que hace útil la bandeja ANTES de Siigo. Lo que estas pruebas vigilan:
que solo salga lo aprobado, que el archivo cuadre, y que un mes sin aprobados
deje un archivo honesto y no un error.
"""
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook

from contab_core import asiento, documentos, export, perfil_fiscal, storage, tarifas, xml_ubl

MUESTRAS = Path(__file__).resolve().parents[3] / "docs" / "samples" / "xml"
CUFE = "1" * 92 + "aaaa"


@pytest.fixture
def conn(tmp_path: Path):
    db = tmp_path / "e.db"
    storage.init_empresa_db(db, nit="900181831", dv="4", razon_social="REAL")
    c = storage.connect(db)
    storage.ensure_schema(c)
    perfil_fiscal.guardar_perfil(c, {
        "vigencia_desde": "2026-01-01", "regimen": "ordinario", "responsable_iva": 1,
        "periodicidad_iva": "cuatrimestral", "agente_retefuente": 1, "agente_reteiva": 1})
    tarifas.guardar_uvt(c, 2026, "52374")
    tarifas.guardar_concepto(c, {
        "concepto": "Compras.", "aplica_a": "ambos", "sobre": "base", "tarifa": "2.5",
        "base_minima_uvt": "10", "cuenta_puc": "23654001", "vigencia_desde": "2026-01-01"})
    lectura = xml_ubl.leer(MUESTRAS / "factura-con-iva.xml")
    documentos.guardar(c, lectura.documento, MUESTRAS / "factura-con-iva.xml")
    c.execute("INSERT INTO excel_load(id, source_path, source_name, total_rows, periodo)"
              " VALUES (1,'x.xlsx','x.xlsx',1,'2026-09')")
    c.execute("INSERT INTO excel_row(load_id,row_index,cufe_cude,data_json,periodo)"
              " VALUES (1,0,?,'{\"Grupo\":\"Recibido\"}','2026-09')", (CUFE,))
    c.execute("UPDATE empresa_info SET cuenta_proveedores='220505', cuenta_iva_descontable='240810'")
    c.commit()
    yield c
    c.close()


def _filas(path: str) -> list[list]:
    ws = load_workbook(path).active
    return [[cell.value for cell in row] for row in ws.iter_rows()]


def test_solo_sale_lo_aprobado(conn, tmp_path: Path):
    asiento.generar(conn, "2026-09")
    r = export.export_asientos(conn, "2026-09", tmp_path)
    assert r["asientos"] == 0 and r["lineas"] == 0
    assert Path(r["path"]).exists(), "un mes sin aprobados deja un archivo honesto, no un error"

    a = asiento.listar(conn, "2026-09")[0]
    concepto = tarifas.listar_conceptos(conn)[0]["id"]
    asiento.aprobar(conn, a["id"], cuenta_puc="519530", concepto_retencion_id=concepto)
    r = export.export_asientos(conn, "2026-09", tmp_path)
    assert r["asientos"] == 1 and r["lineas"] == 4
    assert r["path"].endswith("asientos-2026-09.xlsx")


def test_el_archivo_cuadra_y_lo_dice(conn, tmp_path: Path):
    asiento.generar(conn, "2026-09")
    a = asiento.listar(conn, "2026-09")[0]
    concepto = tarifas.listar_conceptos(conn)[0]["id"]
    asiento.aprobar(conn, a["id"], cuenta_puc="519530", concepto_retencion_id=concepto)
    r = export.export_asientos(conn, "2026-09", tmp_path)
    assert r["cuadra"] is True
    assert Decimal(r["total_debito"]) == Decimal(r["total_credito"]) == Decimal("1785000")


def test_una_fila_por_linea_con_las_columnas_que_entiende_siigo(conn, tmp_path: Path):
    asiento.generar(conn, "2026-09")
    a = asiento.listar(conn, "2026-09")[0]
    concepto = tarifas.listar_conceptos(conn)[0]["id"]
    asiento.aprobar(conn, a["id"], cuenta_puc="519530", concepto_retencion_id=concepto)
    filas = _filas(export.export_asientos(conn, "2026-09", tmp_path)["path"])
    assert filas[0] == export.ASIENTO_COLUMNS
    cuentas = {f[2] for f in filas[1:5]}
    assert cuentas == {"519530", "240810", "23654001", "220505"}
    gasto = next(f for f in filas[1:5] if f[2] == "519530")
    assert gasto[4] == 1500000 and gasto[6] == CUFE, "débito y CUFE en cada línea"
