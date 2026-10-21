"""El importador de tarifas — 4.1.

Las muestras se construyen aquí con openpyxl y no se guardan en `docs/samples/`
porque el archivo real trae NIT y razón social de clientes. Lo que sí se copia
del real es **la forma**, que está verificada en `docs/tarifas-structure.md`:
tarifas como fracción con formato de porcentaje, «N/A» por base mínima, y las
dos filas de ReteIVA cuya tarifa es texto.

La prueba que más vale es la primera: leer `0.11` como si fuera el porcentaje da
0,11% en vez de 11% — cien veces menos retención en todas las filas, sin que
nada falle.
"""
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import Workbook

from contab_core import storage, tarifas


@pytest.fixture
def excel(tmp_path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "1. UVT"
    ws.append(["Año", "Valor de la UVT (pesos)", "Norma de referencia"])
    ws.append([2026, 52374, "Resolución DIAN 000238 de 2025"])
    ws.append([2025, 49799, "Resolución DIAN 193 de 2024"])

    ws = wb.create_sheet("2. Retenciones")
    ws.append(["Concepto", "Tarifa (%)", "Base mínima (UVT)", "Base mínima (pesos)",
               "Cuenta PUC donde lo contabiliza", "Aplica a", "Rige desde", "Observaciones"])
    filas = [
        ("Honorarios", 0.11, "N/A", 1, "23651503", "Declarantes", "2026-01-01", ""),
        ("Compras.", 0.025, 10, 0, "23654001", "Declarantes", "2026-01-01", ""),
        ("Compras.", 0.035, 10, 0, "23654004", "No Declarantes", "2026-01-01", ""),
        ("Servicios en general", 0.04, 2, 0, "23652503", "Declarantes y No Declarantes",
         "2026-01-01", ""),
        ("Cuando adquieran servicios gravados.", "15% del IVA", "Bases retención renta",
         "IVA Facturado", "23670103", "Régimen simple", "2026-01-01", ""),
    ]
    for f in filas:
        ws.append(list(f))
    # Lo que hace la trampa: las celdas de tarifa son celdas de PORCENTAJE.
    for r in range(2, 2 + len(filas) - 1):
        ws.cell(r, 2).number_format = "0%"

    destino = tmp_path / "tarifas.xlsx"
    wb.save(destino)
    return destino


@pytest.fixture
def conn(tmp_path: Path):
    db = tmp_path / "e.db"
    storage.init_empresa_db(db, nit="900181831", dv="4", razon_social="REAL")
    c = storage.connect(db)
    storage.ensure_schema(c)
    yield c
    c.close()


# ───────────────────── la trampa del porcentaje ────────────────────────────────

def test_una_celda_de_porcentaje_se_lee_como_porcentaje(excel: Path):
    """0.11 con formato '0%' son 11%, no 0,11%."""
    c = tarifas.cargar(excel)
    hon = next(x for x in c.conceptos if x["concepto"] == "Honorarios")
    assert Decimal(hon["tarifa"]) == Decimal("11.00")


def test_tambien_se_acepta_la_tarifa_escrita_como_numero(tmp_path: Path):
    """El siguiente contador puede escribir 11 en una celda normal."""
    wb = Workbook()
    ws = wb.active
    ws.title = "UVT"
    ws.append(["Año", "Valor"])
    ws.append([2026, 52374])
    ws = wb.create_sheet("Retenciones")
    ws.append(["Concepto", "Tarifa (%)", "Base mínima (UVT)", "Cuenta PUC"])
    ws.append(["Honorarios", 11, "N/A", "23651503"])   # sin formato de porcentaje
    destino = tmp_path / "otro.xlsx"
    wb.save(destino)
    c = tarifas.cargar(destino)
    assert Decimal(c.conceptos[0]["tarifa"]) == Decimal("11")


def test_las_tarifas_son_decimal_no_float(excel: Path):
    c = tarifas.cargar(excel)
    for x in c.conceptos:
        if x["tarifa"]:
            assert isinstance(Decimal(x["tarifa"]), Decimal)


# ───────────────────── lo que el Excel real trae raro ──────────────────────────

def test_na_en_la_base_minima_significa_que_no_hay_minimo(excel: Path):
    """No es un dato que falte: se retiene desde el primer peso."""
    c = tarifas.cargar(excel)
    hon = next(x for x in c.conceptos if x["concepto"] == "Honorarios")
    assert hon["base_minima_uvt"] is None
    assert tarifas.base_minima_de("N/A") is None
    assert tarifas.base_minima_de(10) == Decimal("10")


def test_reteiva_queda_marcada_como_tarifa_sobre_el_iva(excel: Path):
    """Sin esta marca se calcularía sobre la base gravable y saldría mal."""
    c = tarifas.cargar(excel)
    iva = [x for x in c.conceptos if x["sobre"] == tarifas.SOBRE_IVA]
    assert len(iva) == 1
    assert iva[0]["tarifa"] is None
    assert iva[0]["tarifa_texto"] == "15% del IVA"


def test_una_tarifa_que_no_es_numero_se_avisa_y_no_se_pierde(excel: Path):
    c = tarifas.cargar(excel)
    assert any("15% del IVA" in a for a in c.avisos)


def test_el_mismo_concepto_dos_veces_segun_el_tercero(excel: Path):
    """«Compras.» con 2,5% a declarantes y 3,5% a no declarantes."""
    c = tarifas.cargar(excel)
    compras = sorted((x for x in c.conceptos if x["concepto"] == "Compras."),
                     key=lambda x: x["aplica_a"])
    assert [x["aplica_a"] for x in compras] == ["declarantes", "no_declarantes"]
    assert Decimal(compras[0]["tarifa"]) == Decimal("2.500")
    assert Decimal(compras[1]["tarifa"]) == Decimal("3.500")


# ───────────────────────── guardar y consultar ─────────────────────────────────

def test_la_base_minima_en_pesos_se_calcula_con_la_uvt(excel: Path, conn):
    """La columna de pesos del Excel real venía vacía: se calcula aquí."""
    tarifas.guardar(conn, tarifas.cargar(excel))
    compras = next(x for x in tarifas.listar_conceptos(conn)
                   if x["concepto"] == "Compras." and x["aplica_a"] == "declarantes")
    assert tarifas.base_minima_pesos(conn, compras, 2026) == Decimal("523740")
    assert tarifas.base_minima_pesos(conn, compras, 2025) == Decimal("497990")


def test_sin_minimo_no_hay_base_en_pesos(excel: Path, conn):
    tarifas.guardar(conn, tarifas.cargar(excel))
    hon = next(x for x in tarifas.listar_conceptos(conn) if x["concepto"] == "Honorarios")
    assert tarifas.base_minima_pesos(conn, hon, 2026) is None


def test_volver_a_subir_el_archivo_fusiona_y_no_duplica(excel: Path, conn):
    tarifas.guardar(conn, tarifas.cargar(excel))
    antes = len(tarifas.listar_conceptos(conn))
    r = tarifas.guardar(conn, tarifas.cargar(excel))
    assert len(tarifas.listar_conceptos(conn)) == antes
    assert r["nuevos"] == 0 and r["actualizados"] > 0


def test_lo_escrito_a_mano_sobrevive_a_una_importacion(excel: Path, conn):
    """Igual que el PUC: subir el archivo fusiona, nunca borra."""
    tarifas.guardar(conn, tarifas.cargar(excel))
    tarifas.guardar_concepto(conn, {
        "concepto": "Un concepto mío", "aplica_a": "ambos", "sobre": "base",
        "tarifa": "7", "cuenta_puc": "236599", "vigencia_desde": "2026-01-01"})
    tarifas.guardar(conn, tarifas.cargar(excel))
    mios = [x for x in tarifas.listar_conceptos(conn) if x["fuente"] == "manual"]
    assert len(mios) == 1 and mios[0]["concepto"] == "Un concepto mío"


def test_el_ensayo_en_seco_no_escribe_nada(excel: Path, conn):
    r = tarifas.guardar(conn, tarifas.cargar(excel), dry_run=True)
    assert r["dry_run"] is True
    assert tarifas.listar_conceptos(conn) == []
    assert tarifas.listar_uvt(conn) == []


def test_cambiar_una_tarifa_a_mano_no_exige_tocar_codigo(excel: Path, conn):
    """Done-when de la fase: cambiar una tarifa cambia el cálculo del mes."""
    tarifas.guardar(conn, tarifas.cargar(excel))
    hon = next(x for x in tarifas.listar_conceptos(conn) if x["concepto"] == "Honorarios")
    tarifas.guardar_concepto(conn, {**hon, "tarifa": "4"})
    de_nuevo = next(x for x in tarifas.listar_conceptos(conn) if x["id"] == hon["id"])
    assert Decimal(de_nuevo["tarifa"]) == Decimal("4")


def test_un_archivo_sin_las_hojas_lo_dice_en_vez_de_quedarse_callado(tmp_path: Path):
    wb = Workbook()
    wb.active.title = "Otra cosa"
    destino = tmp_path / "vacio.xlsx"
    wb.save(destino)
    c = tarifas.cargar(destino)
    assert c.conceptos == [] and c.uvt == []
    assert any("UVT" in a for a in c.avisos)
    assert any("retenciones" in a for a in c.avisos)
