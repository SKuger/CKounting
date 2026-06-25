"""D-048 — skip V1 (and the faltantes report) for self-emitted rows."""
from pathlib import Path

import pytest
from openpyxl import load_workbook

from contab_core import excel, storage, validations
from contab_core.sidecar import HANDLERS


@pytest.fixture
def empresa_with_load(tmp_path: Path, sample_excel: Path) -> dict:
    """REAL Seguros' NIT is 900181831. In the sample, several rows have
    NIT Emisor = 900183118 (close NIT but different) — those are the rows
    REAL emitted on its behalf. We use 900183118 here so they ARE treated as
    self-emitted by the test fixture."""
    db_path = tmp_path / "empresa.db"
    storage.init_empresa_db(db_path, nit="900183118", dv="9", razon_social="REAL")
    loaded = excel.load_excel(sample_excel)
    conn = storage.connect(db_path)
    try:
        load_id = excel.persist_load(conn, str(loaded.rows and sample_excel), loaded)
    finally:
        conn.close()
    pdfs_folder = tmp_path / "pdfs_vacia"
    pdfs_folder.mkdir()
    storage.set_pdfs_folder(db_path, str(pdfs_folder))
    return {"db_path": db_path, "load_id": load_id, "rows": loaded.rows}


def test_v1_skips_self_emitted(empresa_with_load):
    """run_validations must drop V1 failures whose NIT Emisor matches the
    empresa NIT (900183118 in this fixture)."""
    db_path = empresa_with_load["db_path"]
    load_id = empresa_with_load["load_id"]

    res = HANDLERS["run_validations"]({"db_path": str(db_path), "load_id": load_id})
    # Failures should NOT include any row whose NIT Emisor is 900183118.
    rows = empresa_with_load["rows"]
    for f in res["failures"]:
        row = rows[f["row_index"]]
        emisor = validations._normalize_nit(row.get("NIT Emisor"))
        if "V1" in f["codes"]:
            assert emisor != "900183118", (
                f"Row {f['row_index']} ({row.get('# de factura')}) is self-emitted "
                f"but still has V1 failure"
            )


def test_missing_supports_excludes_self_emitted(empresa_with_load, tmp_path: Path):
    db_path = empresa_with_load["db_path"]
    load_id = empresa_with_load["load_id"]
    out = tmp_path / "faltantes.xlsx"

    HANDLERS["export_missing_supports"]({
        "db_path": str(db_path), "load_id": load_id, "output_path": str(out),
    })
    wb = load_workbook(out, read_only=False, data_only=True)
    sheet = wb.active
    headers = [c.value for c in sheet[1]]
    nit_idx = headers.index("NIT Emisor")
    for row in sheet.iter_rows(min_row=2, values_only=True):
        nit = validations._normalize_nit(row[nit_idx])
        assert nit != "900183118", (
            f"Self-emitted NIT should NOT be in the missing-supports report; got row {row}"
        )
    wb.close()


def test_evaluate_v1_self_nit_helper():
    cufe = "a" * 96
    row = {"CUFE/CUDE": cufe, "NIT Emisor": "900183118"}
    # Without self_nit → fails V1.
    assert validations.evaluate_v1(row, set()) is not None
    # With matching self_nit → skipped.
    assert validations.evaluate_v1(row, set(), self_nit="900183118") is None
    # With matching self_nit but different formatting → still skipped.
    assert validations.evaluate_v1(row, set(), self_nit="900.183.118") is None
    # With NON-matching self_nit → still fails V1.
    assert validations.evaluate_v1(row, set(), self_nit="900111111") is not None


def test_override_excludes_from_no_conformes(empresa_with_load):
    """D-046 cross-check requested by user: marking a V3-failing row as approved
    must remove it from failing_rows and from the failures list."""
    db_path = empresa_with_load["db_path"]
    load_id = empresa_with_load["load_id"]
    rows = empresa_with_load["rows"]

    # Find a row that fails V3 in the baseline run.
    baseline = HANDLERS["run_validations"]({"db_path": str(db_path), "load_id": load_id})
    target_idx = None
    for f in baseline["failures"]:
        if "V3" in f["codes"]:
            target_idx = f["row_index"]
            break
    assert target_idx is not None, "sample should contain at least one V3-failing row"

    # Approve it.
    HANDLERS["toggle_validation_override"]({
        "db_path": str(db_path), "load_id": load_id,
        "row_index": target_idx, "code": "V3", "active": True,
    })

    after = HANDLERS["run_validations"]({"db_path": str(db_path), "load_id": load_id})
    overrode_failure = next(
        (f for f in after["failures"] if f["row_index"] == target_idx),
        None,
    )
    if overrode_failure is not None:
        # If the row still appears, V3 must be GONE from codes (V1 may remain).
        assert "V3" not in overrode_failure["codes"]
    # And the global failing_rows count should be ≤ baseline.
    assert after["failing_rows"] <= baseline["failing_rows"]
