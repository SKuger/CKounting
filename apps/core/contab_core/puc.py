"""Plan Único de Cuentas (PUC) por empresa (D-070).

Lee el dump Excel que exporta el software contable — un preámbulo con el nombre y
NIT de la empresa, luego una tabla de cuentas con encabezado
`Código · Nombre · Categoría · Clase · Relación con · Maneja vencimientos ·
Diferencia fiscal · Activo · Nivel agrupación`, y un pie `Procesado en: …`.

Persistencia = **upsert por `codigo` (fusionar)**: el dump actualiza las cuentas
que ya existen y agrega las nuevas; las cuentas manuales que no vienen en el dump
se conservan. `fuente` distingue 'dump' vs 'manual'; `archived` es un ocultado
suave (nunca se borra por importar).
"""
from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from .storage import transaction

# Alias flexibles de cada columna del encabezado (en minúsculas, sin acentos raros).
_COLUMNS: list[tuple[str, tuple[str, ...]]] = [
    ("codigo", ("código", "codigo", "cuenta", "cuenta contable")),
    ("nombre", ("nombre", "descripción", "descripcion", "nombre cuenta")),
    ("categoria", ("categoría", "categoria")),
    ("clase", ("clase",)),
    ("relacion_con", ("relación con", "relacion con", "relación", "relacion")),
    ("maneja_vencimientos", ("maneja vencimientos", "maneja vencimiento")),
    ("diferencia_fiscal", ("diferencia fiscal",)),
    ("activo", ("activo",)),
    ("nivel_agrupacion", ("nivel agrupación", "nivel agrupacion", "nivel de agrupación")),
]

_NIVEL_LABEL = {1: "Clase", 2: "Grupo", 3: "Cuenta", 4: "Subcuenta", 5: "Auxiliar"}


class PucParseError(ValueError):
    """Error de parseo del Excel de PUC, con mensaje en español para la UI."""


@dataclass
class PucLoaded:
    empresa_nombre: str
    empresa_nit: str            # NIT sin DV, solo dígitos ('' si no viene en el Excel)
    accounts: list[dict[str, Any]]
    source_name: str


def _norm(v: Any) -> str:
    return re.sub(r"\s+", " ", str(v if v is not None else "").strip())


def _norm_header(v: Any) -> str:
    return _norm(v).lower()


def _cell_codigo(v: Any) -> str:
    """El código del PUC es numérico. Devuelve la cadena de dígitos o '' (así el
    pie 'Procesado en: …' y filas de texto quedan descartados)."""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = _norm(v)
    return s if re.fullmatch(r"\d+", s) else ""


def _bool_si_no(v: Any) -> int:
    """'Sí'/'Si'/'S'/'1'/True → 1 ; resto → 0. Vacío → 1 (activo por defecto)."""
    t = _norm(v).lower()
    if t == "":
        return 1
    return 1 if t in ("sí", "si", "s", "1", "true", "verdadero", "x", "activo") else 0


def puc_nivel(codigo: str) -> int:
    """Nivel jerárquico por longitud del código (PUC Colombia): 1=Clase(1 díg),
    2=Grupo(2), 3=Cuenta(4), 4=Subcuenta(6), 5=Auxiliar(8)…"""
    n = len(re.sub(r"\D", "", codigo or ""))
    return 1 if n <= 1 else 1 + n // 2


def nivel_label(nivel: int) -> str:
    return _NIVEL_LABEL.get(nivel, "Auxiliar")


def _extract_nit(text: str) -> str:
    """NIT de la empresa desde una celda del preámbulo. Reconoce el formato con DV
    ('900181831-4', 'NIT: 900.181.831-4') y con etiqueta ('Nit 900181831').
    Devuelve el NIT **sin DV** (solo dígitos), o '' si la celda no parece un NIT
    (así una fecha '20260809' o un id no se confunden con el NIT)."""
    t = str(text or "")
    # NIT con dígito de verificación (guión) — el caso confiable: el DV queda fuera.
    m = re.search(r"(\d[\d.\s]{5,14})\s*-\s*\d\b", t)
    if m:
        return re.sub(r"\D", "", m.group(1))
    # Etiqueta 'NIT ...': tomar los dígitos que la siguen.
    m = re.search(r"nit[\s:.\-]*(\d[\d.\s]{5,14})", t, re.I)
    if m:
        return re.sub(r"\D", "", m.group(1))
    return ""


def load_puc(source_path: str | Path) -> PucLoaded:
    """Parsea el dump Excel → PucLoaded. Lanza PucParseError con mensaje claro."""
    from openpyxl import load_workbook

    path = Path(source_path)
    if not path.is_file():
        raise PucParseError(f"No se encontró el archivo: {path}")

    try:
        wb = load_workbook(filename=str(path), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl no lee .xls antiguo ni archivos corruptos
        raise PucParseError(
            "No pude leer el Excel. Asegúrate de que sea un archivo .xlsx válido "
            "(el formato .xls antiguo no es compatible)."
        ) from exc
    try:
        ws = wb.worksheets[0]
        raw = [list(r) for r in ws.iter_rows(values_only=True)]
    finally:
        wb.close()

    # 1) fila de encabezado = la primera que tenga una celda de "código" Y una de
    #    "nombre" (en cualquier columna). Tolerante a variantes: Código/Codigo/Cuenta
    #    + Nombre/Descripción, según cómo lo exporte cada software contable.
    cod_aliases = set(_COLUMNS[0][1])
    nom_aliases = set(_COLUMNS[1][1])
    header_idx = None
    for i, row in enumerate(raw):
        hn = [_norm_header(c) for c in row]
        if any(h in cod_aliases for h in hn) and any(h in nom_aliases for h in hn):
            header_idx = i
            break
    if header_idx is None:
        raise PucParseError(
            "No reconocí la tabla de cuentas en el Excel. Debe tener una fila de "
            "encabezado con columnas 'Código' y 'Nombre' (así lo exporta el software "
            "contable). Revisa que sea el archivo de cuentas contables correcto y que "
            "esté en formato .xlsx."
        )
    header = [_norm_header(c) for c in raw[header_idx]]

    # 2) empresa + NIT del preámbulo (filas antes del encabezado).
    empresa_nombre = ""
    empresa_nit = ""
    for row in raw[:header_idx]:
        for cell in row:
            t = _norm(cell)
            if not t:
                continue
            nit = _extract_nit(t)
            if nit:
                if not empresa_nit:
                    empresa_nit = nit
                continue
            if (not empresa_nombre and any(ch.isalpha() for ch in t)
                    and not t.lower().startswith(("cuentas", "nit"))):
                empresa_nombre = t

    # 3) mapa columna→índice por nombre flexible.
    col_idx: dict[str, int] = {}
    for key, aliases in _COLUMNS:
        for j, h in enumerate(header):
            if h in aliases:
                col_idx[key] = j
                break
    if "codigo" not in col_idx or "nombre" not in col_idx:
        raise PucParseError("El PUC debe tener al menos columnas 'Código' y 'Nombre'.")

    def get(row: list[Any], key: str) -> str:
        j = col_idx.get(key)
        return _norm(row[j]) if j is not None and j < len(row) else ""

    # 4) filas de datos (hasta el pie / filas sin código numérico).
    accounts: list[dict[str, Any]] = []
    seen: set[str] = set()
    for row in raw[header_idx + 1:]:
        codigo = _cell_codigo(row[col_idx["codigo"]]) if col_idx["codigo"] < len(row) else ""
        if not codigo or codigo in seen:
            continue
        seen.add(codigo)
        accounts.append({
            "codigo": codigo,
            "nombre": get(row, "nombre"),
            "categoria": get(row, "categoria") or None,
            "clase": get(row, "clase") or None,
            "relacion_con": get(row, "relacion_con") or None,
            "maneja_vencimientos": get(row, "maneja_vencimientos") or None,
            "diferencia_fiscal": get(row, "diferencia_fiscal") or None,
            "activo": _bool_si_no(get(row, "activo")),
            "nivel_agrupacion": get(row, "nivel_agrupacion") or None,
            "nivel": puc_nivel(codigo),
        })
    if not accounts:
        raise PucParseError("El Excel no tiene cuentas debajo del encabezado.")

    return PucLoaded(
        empresa_nombre=empresa_nombre,
        empresa_nit=empresa_nit,
        accounts=accounts,
        source_name=path.name,
    )


# ---------- Persistencia (upsert por código = fusionar) ----------

_UPSERT_FIELDS = (
    "nombre", "categoria", "clase", "relacion_con", "maneja_vencimientos",
    "diferencia_fiscal", "activo", "nivel_agrupacion", "nivel",
)


def persist_puc(conn, loaded: PucLoaded, *, dry_run: bool = False) -> dict[str, Any]:
    """Upsert por `codigo`. Devuelve stats. `dry_run=True` no escribe (solo cuenta
    cuántas se agregarían/actualizarían, para la vista previa)."""
    existing = {r["codigo"] for r in conn.execute("SELECT codigo FROM puc_account").fetchall()}
    to_insert = [a for a in loaded.accounts if a["codigo"] not in existing]
    to_update = [a for a in loaded.accounts if a["codigo"] in existing]

    stats: dict[str, Any] = {
        "total_rows": len(loaded.accounts),
        "inserted": len(to_insert),
        "updated": len(to_update),
        "empresa_nombre": loaded.empresa_nombre,
        "empresa_nit": loaded.empresa_nit,
        "source_name": loaded.source_name,
        "dry_run": dry_run,
        "sample_new": [{"codigo": a["codigo"], "nombre": a["nombre"]} for a in to_insert[:10]],
    }
    if dry_run:
        return stats

    with transaction(conn):
        for a in loaded.accounts:
            if a["codigo"] in existing:
                conn.execute(
                    "UPDATE puc_account SET nombre=?, categoria=?, clase=?, relacion_con=?, "
                    "maneja_vencimientos=?, diferencia_fiscal=?, activo=?, nivel_agrupacion=?, "
                    "nivel=?, fuente='dump', updated_at=datetime('now') WHERE codigo=?",
                    (a["nombre"], a["categoria"], a["clase"], a["relacion_con"],
                     a["maneja_vencimientos"], a["diferencia_fiscal"], a["activo"],
                     a["nivel_agrupacion"], a["nivel"], a["codigo"]),
                )
            else:
                conn.execute(
                    "INSERT INTO puc_account (codigo, nombre, categoria, clase, relacion_con, "
                    "maneja_vencimientos, diferencia_fiscal, activo, nivel_agrupacion, nivel, fuente) "
                    "VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 'dump')",
                    (a["codigo"], a["nombre"], a["categoria"], a["clase"], a["relacion_con"],
                     a["maneja_vencimientos"], a["diferencia_fiscal"], a["activo"],
                     a["nivel_agrupacion"], a["nivel"]),
                )
        conn.execute(
            "INSERT INTO puc_import (source_name, excel_nit, excel_razon, total_rows, inserted, updated) "
            "VALUES (?, ?, ?, ?, ?, ?)",
            (loaded.source_name, loaded.empresa_nit, loaded.empresa_nombre,
             stats["total_rows"], stats["inserted"], stats["updated"]),
        )
    return stats


# ---------- Lectura / CRUD manual ----------

def list_accounts(conn, *, include_archived: bool = True) -> list[dict[str, Any]]:
    q = "SELECT * FROM puc_account"
    if not include_archived:
        q += " WHERE archived = 0"
    q += " ORDER BY codigo"
    return [dict(r) for r in conn.execute(q).fetchall()]


def last_import(conn) -> dict[str, Any] | None:
    r = conn.execute("SELECT * FROM puc_import ORDER BY id DESC LIMIT 1").fetchone()
    return dict(r) if r else None


def counts(conn) -> dict[str, int]:
    row = conn.execute(
        "SELECT COUNT(*) total, "
        "SUM(CASE WHEN archived=0 THEN 1 ELSE 0 END) activas, "
        "SUM(CASE WHEN archived=1 THEN 1 ELSE 0 END) archivadas, "
        "SUM(CASE WHEN fuente='manual' THEN 1 ELSE 0 END) manuales "
        "FROM puc_account"
    ).fetchone()
    return {k: (row[k] or 0) for k in ("total", "activas", "archivadas", "manuales")}


def save_account(conn, data: dict[str, Any], *, account_id: int | None = None) -> dict[str, Any]:
    """Crea (account_id None) o edita una cuenta. El código es único; editarlo
    revalida la unicidad. Las creadas a mano quedan con fuente='manual'."""
    codigo = re.sub(r"\D", "", str(data.get("codigo") or ""))
    if not codigo:
        raise PucParseError("El código de la cuenta es obligatorio y debe ser numérico.")
    nombre = _norm(data.get("nombre"))
    if not nombre:
        raise PucParseError("El nombre de la cuenta es obligatorio.")

    vals = {
        "codigo": codigo,
        "nombre": nombre,
        "categoria": _norm(data.get("categoria")) or None,
        "clase": _norm(data.get("clase")) or None,
        "relacion_con": _norm(data.get("relacion_con")) or None,
        "maneja_vencimientos": _norm(data.get("maneja_vencimientos")) or None,
        "diferencia_fiscal": _norm(data.get("diferencia_fiscal")) or None,
        "activo": 1 if data.get("activo", True) else 0,
        "nivel_agrupacion": _norm(data.get("nivel_agrupacion")) or None,
        "nivel": puc_nivel(codigo),
    }

    with transaction(conn):
        if account_id is None:
            clash = conn.execute("SELECT 1 FROM puc_account WHERE codigo=?", (codigo,)).fetchone()
        else:
            clash = conn.execute(
                "SELECT 1 FROM puc_account WHERE codigo=? AND id<>?", (codigo, account_id)
            ).fetchone()
        if clash:
            raise PucParseError(f"Ya existe una cuenta con el código {codigo}.")

        if account_id is None:
            cur = conn.execute(
                "INSERT INTO puc_account (codigo, nombre, categoria, clase, relacion_con, "
                "maneja_vencimientos, diferencia_fiscal, activo, nivel_agrupacion, nivel, fuente) "
                "VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 'manual')",
                tuple(vals[k] for k in ("codigo", "nombre", "categoria", "clase", "relacion_con",
                                        "maneja_vencimientos", "diferencia_fiscal", "activo",
                                        "nivel_agrupacion", "nivel")),
            )
            account_id = cur.lastrowid
        else:
            n = conn.execute(
                "UPDATE puc_account SET codigo=?, nombre=?, categoria=?, clase=?, relacion_con=?, "
                "maneja_vencimientos=?, diferencia_fiscal=?, activo=?, nivel_agrupacion=?, nivel=?, "
                "updated_at=datetime('now') WHERE id=?",
                (*tuple(vals[k] for k in ("codigo", "nombre", "categoria", "clase", "relacion_con", "maneja_vencimientos", "diferencia_fiscal", "activo", "nivel_agrupacion", "nivel")), account_id),
            ).rowcount
            if not n:
                raise PucParseError("La cuenta que intentas editar ya no existe.")

    r = conn.execute("SELECT * FROM puc_account WHERE id=?", (account_id,)).fetchone()
    return dict(r)


def set_archived(conn, account_id: int, archived: bool) -> None:
    with transaction(conn):
        n = conn.execute(
            "UPDATE puc_account SET archived=?, updated_at=datetime('now') WHERE id=?",
            (1 if archived else 0, account_id),
        ).rowcount
    if not n:
        raise PucParseError("La cuenta no existe.")


def delete_account(conn, account_id: int) -> None:
    with transaction(conn):
        conn.execute("DELETE FROM puc_account WHERE id=?", (account_id,))


def export_puc(conn, output_path: str | Path, *, include_archived: bool = False) -> dict[str, Any]:
    """Exporta el PUC a .xlsx con las columnas originales. Devuelve {path, rows_written}."""
    import xlsxwriter

    out = Path(output_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    accounts = list_accounts(conn, include_archived=include_archived)

    wb = xlsxwriter.Workbook(str(out))
    ws = wb.add_worksheet("PUC")
    header_fmt = wb.add_format({"bold": True, "bg_color": "#EFEFEF", "border": 1})
    cell_fmt = wb.add_format({"border": 1})

    cols = ["Código", "Nombre", "Categoría", "Clase", "Relación con",
            "Maneja vencimientos", "Diferencia fiscal", "Activo", "Nivel agrupación"]
    for c, name in enumerate(cols):
        ws.write(0, c, name, header_fmt)
    ws.set_column(0, 0, 14)
    ws.set_column(1, 1, 44)
    ws.set_column(2, 8, 18)

    keys = ["codigo", "nombre", "categoria", "clase", "relacion_con",
            "maneja_vencimientos", "diferencia_fiscal", "activo", "nivel_agrupacion"]
    for r, acc in enumerate(accounts, start=1):
        for c, k in enumerate(keys):
            v = acc.get(k)
            if k == "activo":
                v = "Sí" if v else "No"
            ws.write(r, c, "" if v is None else str(v), cell_fmt)

    ws.freeze_panes(1, 0)
    wb.close()
    return {"path": str(out), "rows_written": len(accounts)}
