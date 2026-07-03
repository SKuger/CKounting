"""Consulta de Inconsistencias DIAN (razón social por NIT+DV).

Scrapea la página pública de MUISCA `DefSelPublicacionesExterna.faces`
(JSF/Mojarra, sin Cloudflare) por **HTTP plano** — sin navegador. Dado un NIT
y su dígito de verificación, devuelve la razón social.

Flujo: GET (trae cookies de sesión + ViewState) → POST (numNit + dv + ViewState)
→ la respuesta ya trae la razón social en el HTML. Solo usa la librería estándar
(`urllib`), así que no agrega dependencias al sidecar.
"""
from __future__ import annotations

import html as _html
import http.cookiejar
import re
import urllib.parse
import urllib.request
from pathlib import Path
from typing import Any

from . import nit as nit_module

_HOST = "https://muisca.dian.gov.co"
_BASE = f"{_HOST}/WebGestionmasiva/DefSelPublicacionesExterna.faces"
_UA = (
    "Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 "
    "(KHTML, like Gecko) Chrome/120 Safari/537.36"
)
_FORM = "vistaSelPublicacionesExterna:formSelPublicacionesExterna"


def _clean(fragment: str) -> str:
    """Strip HTML tags + decode entities + collapse whitespace."""
    text = _html.unescape(re.sub(r"<[^>]+>", " ", fragment))
    return re.sub(r"\s+", " ", text).strip()


def _hidden(html: str, name: str) -> str:
    """Read a hidden input value regardless of attribute order."""
    m = re.search(r'name="' + re.escape(name) + r'"[^>]*\bvalue="([^"]*)"', html, re.S)
    if m:
        return m.group(1)
    m = re.search(r'\bvalue="([^"]*)"[^>]*name="' + re.escape(name) + r'"', html, re.S)
    return m.group(1) if m else ""


def consultar(nit: str, dv: str, *, timeout: int = 25) -> dict[str, Any]:
    """Query one NIT+DV → {nit, dv, razon_social, found, error}.

    Never raises: network/parse problems are returned in `error` so a batch can
    keep going.
    """
    nit_s = str(nit).strip()
    dv_s = str(dv).strip()
    result: dict[str, Any] = {
        "nit": nit_s, "dv": dv_s, "razon_social": "", "found": False, "error": None,
    }
    if not nit_s:
        result["error"] = "NIT vacío"
        return result

    try:
        cj = http.cookiejar.CookieJar()
        opener = urllib.request.build_opener(urllib.request.HTTPCookieProcessor(cj))
        opener.addheaders = [
            ("User-Agent", _UA),
            ("Accept", "text/html,application/xhtml+xml"),
        ]

        # 1) GET: establece sesión + ViewState
        page = opener.open(_BASE, timeout=timeout).read().decode("iso-8859-1", "replace")
        viewstate = _hidden(page, "com.sun.faces.VIEW")
        m = re.search(r'<form[^>]*\baction="([^"]*)"', page)
        action = (_HOST + m.group(1)) if m else _BASE

        # 2) POST: campos del formulario + criterios de búsqueda
        data = {
            f"{_FORM}:modoPresentacionSeleccionBO": "pantalla",
            f"{_FORM}:siguienteURL": "",
            f"{_FORM}:modoPresentacionFormBO": "pantalla",
            f"{_FORM}:modoOperacionFormBO": "",
            f"{_FORM}:mantenerCriterios": "",
            f"{_FORM}:numNit": nit_s,
            f"{_FORM}:dv": dv_s,
            f"{_FORM}:btnBuscar.x": "12",
            f"{_FORM}:btnBuscar.y": "8",
            "com.sun.faces.VIEW": viewstate,
            _FORM: _FORM,
        }
        body = urllib.parse.urlencode(data).encode("iso-8859-1")
        req = urllib.request.Request(action, data=body, headers={
            "Content-Type": "application/x-www-form-urlencoded",
            "User-Agent": _UA,
            "Referer": _BASE,
        })
        resp = opener.open(req, timeout=timeout).read().decode("iso-8859-1", "replace")

        # 3) Razón Social = celda que sigue a la etiqueta que termina en "Social"
        m = re.search(r"Social\s*</td>\s*<td[^>]*>(.*?)</td>", resp, re.S | re.I)
        razon = _clean(m.group(1)) if m else ""
        result["razon_social"] = razon
        result["found"] = bool(razon)
        if not razon:
            result["error"] = "Sin razón social (NIT no encontrado o DV incorrecto)"
    except Exception as exc:  # noqa: BLE001 — cualquier fallo de red/parse va a `error`
        result["error"] = f"{exc.__class__.__name__}: {exc}"
    return result


# ---------- Excel entrada / salida ----------

_NIT_HEADERS = {"nit", "numero nit", "número nit", "num nit", "nit tercero"}
_DV_HEADERS = {
    "dv", "digito", "dígito", "digito verificacion", "dígito de verificación",
    "digito de verificacion", "digito verificación", "dv nit", "digito de verificación",
}


def _norm_header(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip().lower())


def _norm_id(value: Any) -> str:
    """Excel devuelve NITs como float ('860002180.0') o con espacios; deja solo dígitos."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return re.sub(r"\D", "", str(value).strip())


def parse_nit_dv(source_path: str | Path) -> dict[str, Any]:
    """Lee un .xlsx con columnas NIT y DV → {rows, detected, total}.

    Detecta los encabezados por nombre (flexible). Si falta el DV de una fila
    pero hay NIT, lo calcula (módulo 11 DIAN). Lanza ValueError con mensaje claro
    si no encuentra la columna NIT.
    """
    from openpyxl import load_workbook

    path = Path(source_path)
    if not path.is_file():
        raise ValueError(f"No se encontró el archivo: {path}")

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.active
        rows_iter = ws.iter_rows(values_only=True)
        header = next(rows_iter, None)
        if header is None:
            raise ValueError("El Excel está vacío.")

        nit_col = dv_col = None
        for idx, cell in enumerate(header):
            h = _norm_header(cell)
            if nit_col is None and (h in _NIT_HEADERS or h == "nit" or h.startswith("nit")):
                nit_col = idx
            elif dv_col is None and (h in _DV_HEADERS or h == "dv" or "verifica" in h):
                dv_col = idx
        if nit_col is None:
            found = ", ".join(str(c) for c in header if c) or "(sin encabezados)"
            raise ValueError(
                "No encontré la columna 'NIT' en el Excel. "
                f"Encabezados leídos: {found}."
            )

        rows: list[dict[str, Any]] = []
        for i, raw in enumerate(rows_iter):
            nit = _norm_id(raw[nit_col]) if nit_col < len(raw) else ""
            dv = ""
            if dv_col is not None and dv_col < len(raw):
                dv = _norm_id(raw[dv_col])
            if not nit:
                continue
            if not dv:  # DV ausente → calcularlo (cortesía)
                try:
                    dv = nit_module.compute_dv(nit)
                except Exception:
                    dv = ""
            rows.append({"row_index": i, "nit": nit, "dv": dv})
    finally:
        wb.close()

    return {
        "rows": rows,
        "total": len(rows),
        "detected": {"nit_col": nit_col, "dv_col": dv_col},
    }


def export_inconsistencias(rows: list[dict[str, Any]], output_path: str | Path) -> dict[str, Any]:
    """Escribe un .xlsx con columnas NIT, DV, Razón Social. Devuelve {path, rows_written}."""
    import xlsxwriter

    out = Path(output_path)
    out.parent.mkdir(parents=True, exist_ok=True)

    wb = xlsxwriter.Workbook(str(out))
    ws = wb.add_worksheet("Inconsistencias")
    header_fmt = wb.add_format({"bold": True, "bg_color": "#EFEFEF", "border": 1})
    cell_fmt = wb.add_format({"border": 1})

    cols = ["NIT", "DV", "Razón Social"]
    for c, name in enumerate(cols):
        ws.write(0, c, name, header_fmt)
    ws.set_column(0, 0, 16)
    ws.set_column(1, 1, 6)
    ws.set_column(2, 2, 48)

    for r, row in enumerate(rows, start=1):
        ws.write(r, 0, str(row.get("nit", "")), cell_fmt)
        ws.write(r, 1, str(row.get("dv", "")), cell_fmt)
        ws.write(r, 2, str(row.get("razon_social", "") or ""), cell_fmt)

    ws.freeze_panes(1, 0)
    wb.close()
    return {"path": str(out), "rows_written": len(rows)}
